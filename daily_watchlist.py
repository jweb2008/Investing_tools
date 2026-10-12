"""
Daily Scanner Watchlist. Runs the evening scan and opens the results in Excel.

  py daily_watchlist.py            full run: (weekly universe rebuild) + scan + watchlist
  py daily_watchlist.py --no-scan  rebuild the watchlist from the latest scan log only

Scheduled by run_daily_scan.bat at 5 PM on weekdays (see README in the
commit message / chat for the Task Scheduler command).

Each run
  1. Skips weekends and NYSE holidays (no new data).
  2. Rebuilds the universe if it is more than 6 days old.
  3. Runs  py scan.py --universe  (which also updates the scan log).
  4. Writes  Watchlists/Daily Scanner Watchlist YYYY-MM-DD.xlsx  dated for the
     NEXT trading day, plus a .txt symbol list for thinkorswim.
  5. Copies the symbols to the clipboard and opens the workbook.

thinkorswim import: Watchlist menu > Import > "Paste symbols from clipboard"
(the symbols are already on the clipboard), or load the .txt file.

If anything fails, a Windows message box says what happened, so a silent
failure never looks like "no setups today".
Read-only: no orders, no account data.
"""
import argparse
import datetime as dt
import os
import subprocess
import sys
import time
from pathlib import Path

import pandas as pd

HERE = Path(__file__).parent
DATA = HERE / "data"
OUT_DIR = HERE / "Watchlists"
LOG_DIR = HERE / "logs"
UNIVERSE = DATA / "universe.txt"
SCAN_LOG = DATA / "scan_log.csv"
UNIVERSE_MAX_AGE_DAYS = 6
from earnings import DANGER_DAYS, WARN_DAYS  # noqa: E402

# NYSE full-day closures. Extend each year.
NYSE_HOLIDAYS = {dt.date.fromisoformat(d) for d in [
    "2026-01-01", "2026-01-19", "2026-02-16", "2026-04-03", "2026-05-25", "2026-06-19",
    "2026-07-03", "2026-09-07", "2026-11-26", "2026-12-25",
    "2027-01-01", "2027-01-18", "2027-02-15", "2027-03-26", "2027-05-31", "2027-06-18",
    "2027-07-05", "2027-09-06", "2027-11-25", "2027-12-24",
]}


def is_trading_day(d):
    return d.weekday() < 5 and d not in NYSE_HOLIDAYS


def next_trading_day(d):
    d += dt.timedelta(days=1)
    while not is_trading_day(d):
        d += dt.timedelta(days=1)
    return d


def today_et():
    return pd.Timestamp.now(tz="America/New_York").date()


def alert(title, msg):
    """Pop-up on Windows; plain print elsewhere."""
    print(f"{title}: {msg}")
    if os.name == "nt":
        try:
            import ctypes
            ctypes.windll.user32.MessageBoxW(0, msg, title, 0x40030)  # icon + topmost
        except Exception:
            pass


def run_step(args, logfile):
    with open(logfile, "a", encoding="utf-8") as f:
        f.write(f"\n===== {dt.datetime.now():%Y-%m-%d %H:%M}  py {' '.join(args)}\n")
        f.flush()
        p = subprocess.run([sys.executable, "-u", *args], cwd=HERE, stdout=f,
                           stderr=subprocess.STDOUT, text=True)
    return p.returncode


def last_lines(logfile, n=6):
    try:
        lines = [l for l in logfile.read_text(encoding="utf-8").splitlines() if l.strip()]
        return "\n".join(lines[-n:])
    except OSError:
        return ""


def tos_symbol(s):
    return s.replace(".", "/")   # BRK.B -> BRK/B


# ------------------------------------------------------------------ workbook
COLS = [  # (header, log column, kind)
    ("Symbol", "Symbol", "text"), ("Status", "Status", "text"), ("Score", "Score", "int"),
    ("Checks", "Checks", "int"), ("Sector", "Sector", "text"),
    ("Close", "Close", "price"), ("Pivot", "Pivot", "price"), ("% Below Pivot", "PctBelowPivot", "pct"),
    ("Entry", "Entry", "price"), ("Stop", "Stop", "price"), ("Risk %", "RiskPct", "pct"),
    ("Target", "Target", "price"), ("Target %", "TargetPct", "pct"), ("R:R", "RewardRisk", "ratio"),
    ("Resistance", "Resistance", "price"), ("Earnings", "EarningsDate", "date"),
    ("Days to Earn.", "DaysToEarnings", "int"), ("Squeeze", "BBSqueeze", "check"),
    ("ATR Contract", "ATRContract", "check"), ("Vol Dry-up", "VolDryup", "check"),
]


def write_workbook(path, rows, near, scan_date, next_day, market):
    from openpyxl import Workbook
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    base = Font(name="Arial", size=10)
    wb = Workbook()

    def sheet(ws, df, title, subtitle, table_name):
        ws.sheet_view.showGridLines = False
        ws["A1"] = title
        ws["A1"].font = Font(name="Arial", size=16, bold=True)
        ws["A2"] = subtitle
        ws["A2"].font = Font(name="Arial", size=10, italic=True, color="555555")
        hdr = 4
        for c, (h, _, _) in enumerate(COLS, 1):
            cell = ws.cell(row=hdr, column=c, value=h)
            cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", start_color="1F4E78")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for i, rec in enumerate(df.to_dict("records"), hdr + 1):
            for c, (_, key, kind) in enumerate(COLS, 1):
                v = rec.get(key)
                if v is None or (isinstance(v, float) and v != v):
                    v = None
                elif kind == "pct":
                    v = float(v) / 100
                elif kind == "check":
                    v = "✓" if v else ""
                elif kind == "int":
                    v = int(v)
                elif kind == "date":
                    v = pd.Timestamp(v).to_pydatetime()
                elif kind == "text" and key == "Symbol":
                    v = str(v)
                cell = ws.cell(row=i, column=c, value=v)
                cell.font = base
                cell.number_format = {"price": "#,##0.00", "pct": "0.0%", "ratio": "0.0",
                                      "int": "0", "date": "ddd m/d"}.get(kind, "General")
                if kind in ("check", "int", "date"):
                    cell.alignment = Alignment(horizontal="center")
        last = hdr + max(len(df), 1)
        ref = f"A{hdr}:{get_column_letter(len(COLS))}{last}"
        t = Table(displayName=table_name, ref=ref)
        t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(t)
        if len(df):
            rr = get_column_letter([k for _, k, _ in COLS].index("RewardRisk") + 1)
            st = get_column_letter([k for _, k, _ in COLS].index("Status") + 1)
            body = f"A{hdr + 1}:{get_column_letter(len(COLS))}{last}"
            ws.conditional_formatting.add(body, FormulaRule(
                formula=[f'${st}{hdr + 1}="Breakout"'],
                fill=PatternFill("solid", start_color="C6EFCE"), font=Font(bold=True)))
            ed = get_column_letter([k for _, k, _ in COLS].index("EarningsDate") + 1)
            dd = get_column_letter([k for _, k, _ in COLS].index("DaysToEarnings") + 1)
            rng = f"{ed}{hdr + 1}:{dd}{last}"
            ws.conditional_formatting.add(rng, FormulaRule(
                formula=[f"AND(ISNUMBER(${dd}{hdr + 1}),${dd}{hdr + 1}<={DANGER_DAYS})"],
                fill=PatternFill("solid", start_color="F8CBAD"), font=Font(bold=True, color="9C0006"),
                stopIfTrue=True))
            ws.conditional_formatting.add(rng, FormulaRule(
                formula=[f"AND(ISNUMBER(${dd}{hdr + 1}),${dd}{hdr + 1}<={WARN_DAYS})"],
                fill=PatternFill("solid", start_color="FFEB9C")))
            ws.conditional_formatting.add(f"{rr}{hdr + 1}:{rr}{last}", FormulaRule(
                formula=[f"AND(ISNUMBER({rr}{hdr + 1}),{rr}{hdr + 1}<2)"],
                font=Font(color="C00000")))
        widths = {"Symbol": 9, "Status": 10, "Sector": 22}
        for c, (h, _, _) in enumerate(COLS, 1):
            ws.column_dimensions[get_column_letter(c)].width = widths.get(h, 11)
        ws.row_dimensions[hdr].height = 28
        ws.freeze_panes = f"B{hdr + 1}"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_title_rows = f"{hdr}:{hdr}"

    word = {True: "above", False: "BELOW", None: "n/a vs"}
    mkt = f"SPY {word.get(market.get('SPY'))} 50-day, QQQ {word.get(market.get('QQQ'))} 50-day"
    nd = f"{next_day:%a} {next_day.month}/{next_day.day}/{next_day:%y}"
    sd = f"{scan_date.month}/{scan_date.day}/{scan_date:%y}"
    n_brk = int((rows["Status"] == "Breakout").sum())
    n_set = int((rows["Status"] == "Setup").sum())
    n_earn = int((pd.to_numeric(rows.get("DaysToEarnings"), errors="coerce") <= WARN_DAYS).sum()) \
        if "DaysToEarnings" in rows else 0

    ws = wb.active
    ws.title = "Watchlist"
    sheet(ws, rows, f"Daily Scanner Watchlist for {nd}",
          f"From the {sd} close.  {n_set} setups, {n_brk} breakout(s), {n_earn} with earnings "
          f"within {WARN_DAYS} days.  Market: {mkt}.  "
          "Shortlist for chart review, not buy signals.", "Watchlist")
    sheet(wb.create_sheet("Near Misses"), near, f"Near Misses for {nd}",
          "Uptrend and near pivot, but fewer contraction checks than required. "
          "Logged to test the cutoff; lower priority.", "NearMisses")

    g = wb.create_sheet("How to Read")
    notes = [
        ("Daily Scanner Watchlist", True),
        ("Setup: uptrend (close > 50 > 200 SMA), within 5% below the pivot, and 2+ of 3 "
         "contraction checks (Squeeze, ATR Contract, Vol Dry-up).", False),
        ("Breakout (green): in an uptrend, closed above the pivot on 1.5x+ average volume.", False),
        ("Pivot: highest high of the prior 30 sessions. Entry is the pivot for setups, the close "
         "for breakouts.", False),
        ("Stop: tighter of the 30-day base low and 2x ATR below entry, never closer than 1x ATR.", False),
        ("Target: lower of the measured move (pivot + base depth) and the nearest overhead "
         "swing high. Blank Resistance = no swing high above entry in the last year.", False),
        ("R:R in red = reward-to-risk under 2:1 to the conservative target.", False),
        (f"Earnings: red = within {DANGER_DAYS} days, yellow = within {WARN_DAYS} days of the scan. "
         "An earnings gap can jump straight through a stop. Dates come from Yahoo Finance and "
         "can be estimates until the company confirms.", False),
        ("All levels are planning references, not predictions.", False),
        ("thinkorswim: Watchlist menu > Import > Paste symbols from clipboard. The run already "
         "copied the symbols; the .txt file next to this workbook has the same list.", False),
    ]
    for i, (text, bold) in enumerate(notes, 1):
        g.cell(row=i, column=1, value=text).font = Font(name="Arial", size=11 if bold else 10, bold=bold)
    g.column_dimensions["A"].width = 120
    g.sheet_view.showGridLines = False

    wb.save(path)


# --------------------------------------------------------------------- main
def build_watchlist(open_it=True):
    if not SCAN_LOG.exists():
        raise RuntimeError("No scan log yet. Run: py scan.py --universe")
    log = pd.read_csv(SCAN_LOG)
    log["ScanDate"] = pd.to_datetime(log["ScanDate"]).dt.date
    scan_date = log["ScanDate"].max()
    day = log[log["ScanDate"] == scan_date].copy()
    if "Target" in day and "Entry" in day:
        day["TargetPct"] = (day["Target"] - day["Entry"]) / day["Entry"] * 100
    else:
        day["TargetPct"] = None
    order = {"Breakout": 0, "Setup": 1}
    rows = (day[day["Status"].isin(order)]
            .assign(_o=lambda d: d["Status"].map(order))
            .sort_values(["_o", "Score", "PctBelowPivot"], ascending=[True, False, True]))
    near = day[day["Status"] == "Near miss"].sort_values(["Score", "PctBelowPivot"],
                                                        ascending=[False, True])
    m = day.iloc[0] if len(day) else {}
    market = {"SPY": _bool(m.get("SPYAbove50")), "QQQ": _bool(m.get("QQQAbove50"))}

    next_day = next_trading_day(scan_date)
    OUT_DIR.mkdir(exist_ok=True)
    stem = f"Daily Scanner Watchlist {next_day.isoformat()}"
    xlsx, txt = OUT_DIR / f"{stem}.xlsx", OUT_DIR / f"{stem}.txt"
    try:
        write_workbook(xlsx, rows, near, scan_date, next_day, market)
    except PermissionError:
        xlsx = OUT_DIR / f"{stem} ({dt.datetime.now():%H%M}).xlsx"   # earlier copy is open
        write_workbook(xlsx, rows, near, scan_date, next_day, market)
    symbols = [tos_symbol(s) for s in rows["Symbol"]]
    txt.write_text("\n".join(symbols) + "\n", encoding="ascii")

    if os.name == "nt":
        try:
            subprocess.run("clip", input="\n".join(symbols), text=True, shell=True, check=True)
        except Exception:
            pass
        if open_it:
            os.startfile(xlsx)
    print(f"Watchlist for {next_day}: {len(rows)} symbols -> {xlsx.name}")
    return scan_date, xlsx


def _bool(v):
    if isinstance(v, str):
        return {"true": True, "false": False}.get(v.strip().lower())
    if v is None or (isinstance(v, float) and v != v):
        return None
    return bool(v)


def main():
    ap = argparse.ArgumentParser(description="Evening scan + Daily Scanner Watchlist")
    ap.add_argument("--no-scan", action="store_true", help="only rebuild the watchlist from the log")
    ap.add_argument("--force", action="store_true", help="run even on a weekend or holiday")
    args = ap.parse_args()

    today = today_et()
    if not args.no_scan:
        if not is_trading_day(today) and not args.force:
            print(f"{today}: market closed today, nothing to scan.")
            return
        LOG_DIR.mkdir(exist_ok=True)
        logfile = LOG_DIR / f"daily_{today.isoformat()}.txt"

        stale = (not UNIVERSE.exists()
                 or (time.time() - UNIVERSE.stat().st_mtime) / 86400 > UNIVERSE_MAX_AGE_DAYS)
        if stale and run_step(["build_universe.py"], logfile) != 0 and not UNIVERSE.exists():
            alert("Daily scan FAILED", "Could not build the universe.\n\n" + last_lines(logfile))
            sys.exit(1)

        if run_step(["scan.py", "--universe"], logfile) != 0:
            alert("Daily scan FAILED", last_lines(logfile)
                  + "\n\nIf this says the Schwab login expired, run: py auth_setup.py")
            sys.exit(1)

    try:
        scan_date, _ = build_watchlist()
    except Exception as e:
        alert("Daily watchlist FAILED", str(e))
        sys.exit(1)
    if not args.no_scan and scan_date != today:
        alert("Daily scan: stale data",
              f"The newest scan in the log is from {scan_date}, not today ({today}). "
              f"Check logs/daily_{today.isoformat()}.txt")


if __name__ == "__main__":
    main()
