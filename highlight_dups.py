"""
Highlights likely duplicate rows in red so you can delete them by hand.

  py highlight_dups.py            highlight, save as the next version, open it
  py highlight_dups.py --list     just print what it found, change nothing
  py highlight_dups.py --blanks   color blank cells to fill (orange) and rows to delete (red)
  py highlight_dups.py --blanks --list

A row is flagged when it has no Sync ID (entered by hand) and the script's own
rows (with Sync IDs) for the same contract already add up to the same trade:
  - same Closed Date, same total quantity, or
  - same Opened Date, same total quantity, or
  - all of that contract's synced rows between its open and close dates, or
  - same Closed Date across all accounts (a hand row that combined accounts).

Rows with Sync IDs are never flagged. Nothing is deleted; you decide.
Close the tracker in Excel before running.
"""
import os
import re
import sys
from collections import defaultdict

import pandas as pd

import update_tracker as U

YELLOW = 65535  # Excel color code
ORANGE = 255 + 204 * 256 + 153 * 65536  # light orange
DELETE_RED = 255 + 150 * 256 + 150 * 65536  # light red: rows to delete

BUY = ("Buy Quantity", "Purchase PRICE PER SHARE", "Fees & Comm")
SELL = ("Sell Quantity", "Sell Price Per Share", "Sale Fees & Comm")
ALWAYS = ("Account", "Opened Date", "Stock Or Option", "Action", "SYMBOL", "Open/Closed")


def num(v):
    try:
        f = abs(float(v))
    except (TypeError, ValueError):
        return None
    return None if f != f else f  # NaN (blank cell) counts as empty


def qty(r):
    for c in ("Buy Quantity", "Sell Quantity"):
        q = num(r[c])
        if q:
            return q
    return None


def day(v):
    return None if U.is_blank(v) or pd.isna(pd.to_datetime(v, errors="coerce")) \
        else pd.Timestamp(v).normalize()


def read_values(path):
    """Table rows with computed values (formulas like =10+5+5 read as 20)."""
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = next((s for s in wb.worksheets if U.TABLE in s.tables), None)
    if ws is None:
        sys.exit(f"Couldn't find the table {U.TABLE} in {path.name}.")
    ref = ws.tables[U.TABLE].ref
    (c1, r1), (c2, r2) = [openpyxl.utils.cell.coordinate_to_tuple(x)[::-1] for x in ref.split(":")]
    headers = [ws.cell(r1, c).value for c in range(c1, c2 + 1)]
    rows = []
    for r in range(r1 + 1, r2 + 1):
        d = {h: ws.cell(r, c1 + i).value for i, h in enumerate(headers)}
        d["_row"] = r
        rows.append(d)
    return pd.DataFrame(rows), ws.title, c1, c2


def find_dups(df):
    df = df[df["SYMBOL"].map(lambda v: not U.is_blank(v))].copy()
    # adjusted contracts (EOSE1, UP1) are the same contract as EOSE, UP for this check
    df["_sym"] = df["SYMBOL"].map(U.norm_sym).map(
        lambda s: re.sub(r"^([A-Z]+)\d( \d{2}/)", r"\1\2", s))
    df["_acct"] = df["Account"].astype(str)
    df["_q"] = df.apply(qty, axis=1)
    df["_open"] = df["Opened Date"].map(day)
    df["_close"] = df["Closed Date"].map(day)
    synced = df[df.get(U.SYNC_COL, pd.Series(dtype=object)).map(lambda v: not U.is_blank(v))]
    hand = df[df.get(U.SYNC_COL, pd.Series(dtype=object)).map(U.is_blank)]

    by_contract = defaultdict(list)
    for _, s in synced.iterrows():
        by_contract[(s["_acct"], s["_sym"])].append(s)
    by_symbol = defaultdict(list)
    for _, s in synced.iterrows():
        by_symbol[s["_sym"]].append(s)

    def total(rows):
        return sum(r["_q"] or 0 for r in rows)

    flagged = []
    for _, h in hand.iterrows():
        q = h["_q"]
        if not q:
            continue
        same = by_contract.get((h["_acct"], h["_sym"]), [])
        reason = None
        if same:
            if pd.notna(h["_close"]) and abs(total([s for s in same if s["_close"] == h["_close"]]) - q) < 1e-6:
                reason = "same close date and quantity as synced rows"
            elif pd.notna(h["_open"]) and abs(total([s for s in same if s["_open"] == h["_open"]]) - q) < 1e-6:
                reason = "same open date and quantity as synced rows"
            elif pd.notna(h["_open"]) and pd.notna(h["_close"]):
                inside = [s for s in same if pd.notna(s["_open"]) and pd.notna(s["_close"])
                          and h["_open"] <= s["_open"] and s["_close"] <= h["_close"]]
                if inside and abs(total(inside) - q) < 1e-6:
                    reason = "combines synced rows in the same date range"
        if reason is None and same and pd.notna(h["_open"]) and pd.notna(h["_close"]):
            flipped = [s for s in same if s["_open"] == h["_close"] and s["_close"] == h["_open"]]
            if flipped and abs(total(flipped) - q) < 1e-6:
                reason = "same trade as synced rows, with open and close dates reversed"
        if reason is None and same and pd.isna(h["_open"]) and pd.notna(h["_close"]):
            # an open hedge entered the old way: the sale date typed as Closed Date
            sold = [s for s in same if s["_open"] == h["_close"] and pd.isna(s["_close"])]
            if sold and abs(total(sold) - q) < 1e-6:
                reason = "same hedge as synced rows (sale date was typed as Closed Date)"
        if reason is None and pd.notna(h["_close"]):
            everyone = [s for s in by_symbol.get(h["_sym"], []) if s["_close"] == h["_close"]]
            if len({s["_acct"] for s in everyone}) > 1 and abs(total(everyone) - q) < 1e-6:
                reason = "combines several accounts' synced rows"
        if reason:
            flagged.append((int(h["_row"]), h["_acct"], str(h["SYMBOL"]), reason))
    return sorted(flagged)


def find_blanks(df):
    """Blank cells that should hold something. Blanks that are correct are skipped:
    an open position's closing side (and Closed Date), and the buy side of an
    open sold-to-open hedge (premium on the sell side, not closed yet)."""
    out = []
    for _, r in df.iterrows():
        if U.is_blank(r.get("SYMBOL")):
            continue
        blank = lambda c: c in r and U.is_blank(r[c])
        # empty share placeholders from corporate actions (mergers, rights,
        # spinoffs): no quantities at all and no 2026 open. Safe to delete;
        # the update never re-adds these.
        sid = str(r.get(U.SYNC_COL) or "")
        if (":pre>" in sid and str(r.get("Stock Or Option") or "").strip().lower() == "stock"
                and all(blank(c) or num(r[c]) is None for c in BUY + SELL)):
            out.append((int(r["_row"]), str(r.get("Account")), str(r.get("SYMBOL")), ["DELETE"]))
            continue
        need = [c for c in ALWAYS if c in r]
        status = str(r.get("Open/Closed") or "").strip().lower()
        if status == "closed":
            need += list(BUY) + list(SELL) + ["Closed Date"]
        else:
            has_sell = any(not blank(c) for c in SELL)
            has_buy = any(not blank(c) for c in BUY)
            need += list(SELL) if has_sell and not has_buy else list(BUY)
        cols = [c for c in need if blank(c)]
        if cols:
            out.append((int(r["_row"]), str(r.get("Account")), str(r.get("SYMBOL")), cols))
    return out


BUSY = (-2147418111, -2147417846)  # Excel "call rejected" / "application busy"


def retry(fn, tries=40, wait=0.5):
    """Run one Excel call, waiting and retrying while Excel says it's busy."""
    import time
    import pywintypes
    for i in range(tries):
        try:
            return fn()
        except pywintypes.com_error as e:
            if e.args and e.args[0] in BUSY and i < tries - 1:
                time.sleep(wait)
                continue
            raise


def col_letter(n):
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def batched_addresses(paint):
    """Group cells by color into comma-separated range addresses short enough
    for one Excel call each (Excel caps a range address at 255 characters)."""
    by_color = defaultdict(list)
    for r, ca, cb, color in paint:
        a = f"{col_letter(ca)}{r}" if ca == cb else f"{col_letter(ca)}{r}:{col_letter(cb)}{r}"
        by_color[color].append(a)
    for color, addrs in by_color.items():
        chunk = []
        for a in addrs:
            if chunk and len(",".join(chunk + [a])) > 240:
                yield color, ",".join(chunk)
                chunk = []
            chunk.append(a)
        if chunk:
            yield color, ",".join(chunk)


def save_highlighted(src, sheet, paint, label):
    """paint: list of (row, first_col, last_col, color). Saves as next version."""
    try:
        import win32com.client as win32
    except ImportError:
        sys.exit("Highlighting needs Excel automation. Run: py -m pip install pywin32")
    dst = U.TRACKER_DIR / U.next_version_name()
    xl = win32.DispatchEx("Excel.Application")
    xl.Visible = False
    xl.DisplayAlerts = False
    wb = None
    try:
        wb = retry(lambda: xl.Workbooks.Open(str(src.resolve()), UpdateLinks=0, ReadOnly=True))
        retry(lambda: setattr(xl, "ScreenUpdating", False))
        ws = retry(lambda: wb.Worksheets(sheet))
        for color, addr in batched_addresses(paint):
            retry(lambda: setattr(ws.Range(addr).Interior, "Color", color))
        retry(lambda: wb.SaveAs(str(dst.resolve()), FileFormat=51))
    finally:
        # the file is already saved by this point; closing problems are harmless
        try:
            if wb is not None:
                wb.Close(False)
        except Exception:
            pass
        try:
            xl.Quit()
        except Exception:
            pass
    print(f"Saved {dst.name} with {label}. {src.name} was not changed.")
    if hasattr(os, "startfile"):
        os.startfile(str(dst.resolve()))


def main_blanks(src, df, sheet, c1):
    headers = [h for h in df.columns if not str(h).startswith("_")]
    found = find_blanks(df)
    print(f"Tracker: {src.name}  (sheet '{sheet}')")
    if not found:
        print("No blank cells left to fill.")
        return
    dels = [f for f in found if f[3] == ["DELETE"]]
    fills = [f for f in found if f[3] != ["DELETE"]]
    for r, acct, sym, cols in fills:
        print(f"  row {r:>4}  {acct:<7} {sym:<28} fill: {', '.join(cols)}")
    for r, acct, sym, _ in dels:
        print(f"  row {r:>4}  {acct:<7} {sym:<28} DELETE (empty corporate-action placeholder)")
    n = sum(len(c) for _, _, _, c in fills)
    print(f"{n} blank cells to fill in {len(fills)} rows (orange); {len(dels)} rows to delete (red).")
    if "--list" in sys.argv:
        return
    last = c1 + len(headers) - 1
    paint = [(r, c1 + headers.index(c), c1 + headers.index(c), ORANGE)
             for r, _, _, cols in fills for c in cols]
    paint += [(r, c1, last, DELETE_RED) for r, *_ in dels]
    save_highlighted(src, sheet, paint,
                     f"{n} blank cells in orange and {len(dels)} delete rows in red")


def main():
    src = U.find_latest_tracker()
    df, sheet, c1, c2 = read_values(src)
    if "--blanks" in sys.argv:
        return main_blanks(src, df, sheet, c1)
    flagged = find_dups(df)
    print(f"Tracker: {src.name}  (sheet '{sheet}')")
    if not flagged:
        print("No likely duplicates found.")
        return
    for r, acct, sym, why in flagged:
        print(f"  row {r:>4}  {acct:<7} {sym:<28} {why}")
    print(f"{len(flagged)} rows flagged.")
    if "--list" in sys.argv:
        return

    save_highlighted(src, sheet, [(r, c1, c2, DELETE_RED) for r, *_ in flagged],
                     f"{len(flagged)} duplicate rows highlighted red")


if __name__ == "__main__":
    main()
