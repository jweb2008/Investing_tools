"""
After Action Review (AAR): grades how trades actually turned out, using the
Excel tracker as the source of truth.

  py aar.py            build aar.xlsx from the newest tracker
  py aar.py --tracker "2025-26_Trading_Tracker_10.9.11.xlsx"   use a specific file

Read-only. Reads the tracker (table Table14), fills_cache.csv and
data/scan_log.csv. Never opens the tracker in Excel, never writes to it, and
never talks to Schwab orders. See AAR_DESIGN.md.

Build step 1 (this version)
  - Tracker rows -> trades: one continuous holding per account and symbol.
    Scale-ins and partial exits are rows inside one trade.
  - Trade net % = sum of recorded Net Profit / sum of recorded TOTAL INVESTMENT
    (Total Sale Value for sold-to-open trades, matching the tracker's % CHANGE).
    Nothing is recalculated from prices.
  - Style from your Action column (Swing, Long, Hedge, Day, ...), with
    aar_tags.csv overrides. Sold-to-open trades are always Hedge.
  - Outcome: Win / Scratch / Loss with a symmetric scratch band.
  - Decision ID: the same contract or stock entered on the same day in several
    accounts is one decision.
Build step 2
  - Scan log flags grouped into setup episodes per ticker.
  - Each trade matched to the most recent flag 1 to 5 trading days before its
    entry (options on the underlying). Source: Scanner, Discretionary, or
    Before scan log (entered before the scan log could have flagged it).
  - Flag details carried onto the trade: status, score, check combination,
    raw indicator values, market filter, scanner entry/stop/targets. For
    stocks: Pre/Post-breakout entry and R-multiple against the scanner stop.
  - Flags table: one row per episode, Taken or Passed.
Build step 3
  - Daily bars from data.py (Schwab, read-only), cached in data/aar_bars so a
    finished trade or setup never costs another request.
  - Swing/Day trades: underlying move, best and worst move while held, and
    for scanner trades whether the stop or each target was reached.
    Options: R and entry type measured on the underlying.
  - Options: days to expiry at entry, in buckets.
  - Open swing trades since the scan log began: where price stands vs your
    entry and the scanner's stop and target, in % and R (closing prices).
  - Every setup episode graded on its own: triggered, stop or target first,
    highest target, 10 and 20 session returns, R. Reports by status, check
    combination, score, market filter, and taken vs passed.

  py aar.py --no-prices      skip Schwab entirely (cached bars only)
  py aar.py --source yahoo   price history from Yahoo instead

Output: aar.xlsx with tables Summary, Trades, Decisions, Flags, TradeRows,
Review, RunInfo (Power Query), and aar_report.html, a dashboard of charts
(see aar_report.py) you open in a browser.

Settings (.env, all optional)
  AAR_SCRATCH_STOCK    scratch band for stocks, in %       (default 2)
  AAR_SCRATCH_OPTION   scratch band for options, in %      (default 10)
  AAR_LONG_SWING_DAYS  flag Swing trades held longer than this many trading
                       days on the Review sheet            (default 40)
  AAR_MATCH_DAYS       a trade is a scanner trade if entered 1 to this many
                       trading days after a flag           (default 5)
  AAR_EPISODE_GAP      flags of one ticker no more than this many trading
                       days apart are one setup episode    (default 5)
  AAR_TRIGGER_DAYS     a setup must trade through Entry within this many
                       sessions to count as triggered      (default 5)
  AAR_OUTCOME_DAYS     sessions to follow a triggered setup (default 20)
  AAR_OUTPUT           output file                         (default aar.xlsx)
  TRACKER_DIR, TRACKER_PREFIX, TRACKER_TABLE   same as update_tracker.py
"""
import argparse
import datetime as dt
import os
import re
import shutil
import sys
import tempfile
import time
import warnings
from pathlib import Path

import numpy as np
import pandas as pd
from dotenv import load_dotenv

load_dotenv()
warnings.filterwarnings("ignore", message="Unknown extension is not supported")  # openpyxl, harmless
warnings.filterwarnings("ignore", message=".*httpx.*deprecated")                # authlib, harmless

HERE = Path(__file__).parent
TRACKER_DIR = Path(os.getenv("TRACKER_DIR", str(HERE)))
PREFIX = os.getenv("TRACKER_PREFIX", "2025-26_Trading_Tracker_")
TABLE = os.getenv("TRACKER_TABLE", "Table14")
FILLS = Path(os.getenv("TRADE_LOG_PATH", str(HERE / "trade_log.xlsx"))).with_name("fills_cache.csv")
SCAN_LOG = HERE / "data" / "scan_log.csv"
TAGS = HERE / "aar_tags.csv"
OUTPUT = Path(os.getenv("AAR_OUTPUT", str(HERE / "aar.xlsx")))
REPORT = OUTPUT.with_name("aar_report.html")

SCRATCH = {"Stock": float(os.getenv("AAR_SCRATCH_STOCK", "2")),
           "Option": float(os.getenv("AAR_SCRATCH_OPTION", "10"))}
LONG_SWING_DAYS = int(os.getenv("AAR_LONG_SWING_DAYS", "40"))
MATCH_DAYS = int(os.getenv("AAR_MATCH_DAYS", "5"))        # entry 1..N trading days after a flag
EPISODE_GAP = int(os.getenv("AAR_EPISODE_GAP", "5"))      # flags this close together = one episode
STATUS_RANK = {"Setup": 3, "Breakout": 2, "Near miss": 1}
BARS_DIR = HERE / "data" / "aar_bars"
TRIGGER_DAYS = int(os.getenv("AAR_TRIGGER_DAYS", "5"))    # setup must trigger within N days
OUTCOME_DAYS = int(os.getenv("AAR_OUTCOME_DAYS", "20"))   # follow a triggered setup N days
DTE_BUCKETS = [(14, "0-14 days"), (45, "15-45 days"), (10**6, "46+ days")]
TARGET_COLS = [("Target2R", "2R"), ("TargetMeasured", "Measured move"),
               ("Resistance", "Resistance"), ("TargetATR", "3x ATR"), ("Target3R", "3R")]

# Action column value -> style. Anything not listed keeps its own name and is
# reported on the Review sheet so you can decide where it belongs.
# STC and BTO mark the front or back end of a hedge.
STYLE_MAP = {
    "swing": "Swing",
    "long": "Long term", "long term": "Long term", "leaps": "Long term",
    "hedge": "Hedge",
    "stc": "Hedge", "bto": "Hedge",   # front or back end of a hedge
    "day": "Day",
}

# NYSE full-day closures, so hold periods count trading days.
NYSE_HOLIDAYS = pd.to_datetime([
    "2023-01-02", "2023-01-16", "2023-02-20", "2023-04-07", "2023-05-29", "2023-06-19",
    "2023-07-04", "2023-09-04", "2023-11-23", "2023-12-25",
    "2024-01-01", "2024-01-15", "2024-02-19", "2024-03-29", "2024-05-27", "2024-06-19",
    "2024-07-04", "2024-09-02", "2024-11-28", "2024-12-25",
    "2025-01-01", "2025-01-09", "2025-01-20", "2025-02-17", "2025-04-18", "2025-05-26",
    "2025-06-19", "2025-07-04", "2025-09-01", "2025-11-27", "2025-12-25",
    "2026-01-01", "2026-01-19", "2026-02-16", "2026-04-03", "2026-05-25", "2026-06-19",
    "2026-07-03", "2026-09-07", "2026-11-26", "2026-12-25",
    "2027-01-01", "2027-01-18", "2027-02-15", "2027-03-26", "2027-05-31", "2027-06-18",
    "2027-07-05", "2027-09-06", "2027-11-25", "2027-12-24",
]).values.astype("datetime64[D]")


def trading_days(start, end):
    """Trading days from start to end (same day = 0)."""
    if pd.isna(start) or pd.isna(end):
        return None
    a = np.datetime64(pd.Timestamp(start).date(), "D")
    b = np.datetime64(pd.Timestamp(end).date(), "D")
    return int(np.busday_count(a, b, holidays=NYSE_HOLIDAYS))


# ------------------------------------------------------------------ symbols
def canon_root(root):
    """EOSE1 (adjusted contract) -> EOSE, same as update_tracker.py."""
    r = re.sub(r"\d+$", "", str(root).strip().upper())
    return r or str(root).strip().upper()


def norm_sym(s):
    """Same contract however the date or strike was typed."""
    t = re.sub(r"\s+", " ", str(s or "")).strip().upper()
    m = re.match(r"^(\S+) (\d{1,2})/(\d{1,2})/(\d{2}|\d{4}) ([\d.]+) ([CP])$", t)
    if m:
        root, mm, dd, yy, k, cp = m.groups()
        yy = yy if len(yy) == 4 else "20" + yy
        try:
            k = f"{float(k):.2f}"
        except ValueError:
            pass
        t = f"{canon_root(root)} {int(mm):02d}/{int(dd):02d}/{yy} {k} {cp}"
    return t


def underlying(sym):
    s = str(sym or "").strip()
    return canon_root(s.split(" ")[0]) if s else ""


def option_expiry(sym):
    m = re.match(r"^\S+ (\d{2})/(\d{2})/(\d{4}) [\d.]+ [CP]$", norm_sym(sym))
    return pd.Timestamp(f"{m.group(3)}-{m.group(1)}-{m.group(2)}") if m else None


def blank(v):
    if v is None or v is pd.NA or v is pd.NaT:
        return True
    if isinstance(v, str):
        return not v.strip()
    try:
        return bool(pd.isna(v))
    except (TypeError, ValueError):
        return False


def num(v):
    try:
        return 0.0 if blank(v) else float(v)
    except (TypeError, ValueError):
        return 0.0


# ------------------------------------------------------------------- inputs
def find_latest_tracker():
    files = [p for p in TRACKER_DIR.glob(f"{PREFIX}*.xlsx") if not p.name.startswith("~$")]
    if not files:
        sys.exit(f"No tracker found in {TRACKER_DIR} matching {PREFIX}*.xlsx")
    return max(files, key=lambda p: p.stat().st_mtime)


def read_tracker(path):
    """Table rows with the values Excel last saved (formulas come in as results).
    Reads a temporary copy, so it works even while the tracker is open in Excel."""
    import openpyxl
    fd, tmp = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    try:
        shutil.copy2(path, tmp)
        wb = openpyxl.load_workbook(tmp, data_only=True)
        ws = next((s for s in wb.worksheets if TABLE in s.tables), None)
        if ws is None:
            sys.exit(f"Couldn't find the table {TABLE} in {path.name}. "
                     "If you renamed it, add TRACKER_TABLE=<new name> to .env.")
        cells = list(ws[ws.tables[TABLE].ref])
        wb.close()
    finally:
        os.remove(tmp)
    headers = [c.value for c in cells[0]]
    df = pd.DataFrame([[c.value for c in r] for r in cells[1:]], columns=headers)
    df.insert(0, "ExcelRow", [r[0].row for r in cells[1:]])
    for c in ("Opened Date", "Closed Date"):
        df[c] = pd.to_datetime(df[c], errors="coerce")
    needed = {"Account", "Opened Date", "Stock Or Option", "SYMBOL", "Closed Date",
              "TOTAL INVESTMENT", "Total Sale Value", "Net Profit", "Open/Closed"}
    missing = needed - set(df.columns)
    if missing:
        sys.exit(f"{TABLE} is missing columns: {', '.join(sorted(missing))}")
    for c in ("Action", "Sync ID", "Buy Quantity", "Sell Quantity",
              "Purchase PRICE PER SHARE", "% CHANGE"):
        if c not in df.columns:
            df[c] = None
    keep = ~df[["Account", "SYMBOL", "Opened Date", "Closed Date"]].apply(
        lambda r: all(blank(v) for v in r), axis=1)
    return df[keep].reset_index(drop=True), path.name


def short_open_fills():
    """Fill IDs that opened a position by selling (sold to open)."""
    if not FILLS.exists():
        return set()
    f = pd.read_csv(FILLS, dtype=str, usecols=["FillID", "Side", "PositionEffect"])
    return set(f.loc[(f["Side"] == "SELL") & (f["PositionEffect"] == "OPENING"), "FillID"])


def load_tags():
    """aar_tags.csv: Account, Symbol, OpenedDate (optional), Style."""
    if not TAGS.exists():
        return []
    t = pd.read_csv(TAGS, dtype=str).fillna("")
    out = []
    for r in t.itertuples(index=False):
        d = pd.to_datetime(r.OpenedDate, errors="coerce") if r.OpenedDate.strip() else None
        out.append((r.Account.strip(), norm_sym(r.Symbol), d, r.Style.strip()))
    return out


# ------------------------------------------------------------- rows -> trades
def prepare_rows(trk, short_fills):
    r = trk.copy()
    r["Asset"] = np.where(r["Stock Or Option"].astype(str).str.strip().str.lower()
                          .str.startswith("opt"), "Option", "Stock")
    r["Symbol"] = r["SYMBOL"].map(norm_sym)
    r["Underlying"] = r["SYMBOL"].map(underlying)
    r["Account"] = r["Account"].astype(str).str.strip()
    r["ActionRaw"] = r["Action"].map(lambda v: "" if blank(v) else str(v).strip())

    def is_short(row):
        sync = "" if blank(row["Sync ID"]) else str(row["Sync ID"])
        m = re.match(r"^[^:]+:(.*?)>", sync)
        if m:
            opener = m.group(1)
            if opener in short_fills or re.search(r"seed:\d+:S", opener):
                return True
        buy, sell = num(row["Buy Quantity"]), num(row["Sell Quantity"])
        if buy < 0:                       # old "Hedge" convention
            return True
        return buy == 0 and sell > 0 and num(row["TOTAL INVESTMENT"]) == 0
    r["Short"] = r.apply(is_short, axis=1)
    r["Closed"] = r["Open/Closed"].astype(str).str.strip().str.lower().eq("closed")
    # each row's own style, so a swing lot on top of a long-term core holding
    # in the same account stays a separate trade
    r["RowStyle"] = [("Hedge" if sh else STYLE_MAP.get(a.lower(), a) if a else "")
                     for sh, a in zip(r["Short"], r["ActionRaw"])]
    return r


def chain_trades(rows):
    """Assign TradeID: rows of one account+symbol+style whose holding periods
    touch or overlap form one trade (flat to open to flat)."""
    rows = rows.copy()
    rows["TradeID"] = None
    n = 0
    for _, g in rows.groupby(["Account", "Symbol", "RowStyle"], sort=False):
        g = g.sort_values(["Opened Date", "Closed Date"], na_position="first")
        cur_end, cur_open, cur_id = None, False, None
        for idx, row in g.iterrows():
            o, c = row["Opened Date"], row["Closed Date"]
            if pd.isna(o):                # can't place it in time: its own trade
                n += 1
                rows.at[idx, "TradeID"] = f"T{n:04d}"
                continue
            joins = cur_id is not None and (cur_open or (cur_end is not None and o <= cur_end))
            if not joins:
                n += 1
                cur_id, cur_end, cur_open = f"T{n:04d}", None, False
            rows.at[idx, "TradeID"] = cur_id
            if row["Closed"] and not pd.isna(c):
                cur_end = c if cur_end is None else max(cur_end, c)
            else:
                cur_open = True
    return rows


def style_for(trade_rows, tags):
    first = trade_rows.sort_values("Opened Date", na_position="first").iloc[0]
    acct, sym, opened = first["Account"], first["Symbol"], first["Opened Date"]
    if trade_rows["Short"].any():
        return "Hedge", "sold to open"
    for t_acct, t_sym, t_date, t_style in tags:
        if t_acct.lower() == acct.lower() and t_sym == sym and (
                t_date is None or (not pd.isna(opened) and t_date.date() == opened.date())):
            return t_style, "aar_tags.csv"
    actions = [a for a in trade_rows.sort_values("Opened Date", na_position="first")["ActionRaw"] if a]
    if actions:
        a = actions[0]
        return STYLE_MAP.get(a.lower(), a), "Action column"
    if first["Asset"] == "Option" and not pd.isna(opened):
        exp = option_expiry(sym)
        if exp is not None and (exp - opened).days >= 365:
            return "Long term", "LEAPS (1y+ to expiry)"
    return "Untagged", ""


def outcome(pct, asset):
    if pct is None or pd.isna(pct):
        return ""
    band = SCRATCH[asset]
    return "Win" if pct > band else ("Loss" if pct < -band else "Scratch")


def build_trades(rows, tags):
    out, review = [], []
    for tid, g in rows.groupby("TradeID", sort=False):
        g = g.sort_values(["Opened Date", "Closed Date"], na_position="first")
        first = g.iloc[0]
        style, src = style_for(g, tags)
        closed = bool(g["Closed"].all())
        invest = g["TOTAL INVESTMENT"].map(num).sum()
        sale = g["Total Sale Value"].map(num).sum()
        net = g["Net Profit"].map(num).sum()
        base = invest if invest > 0 else sale
        pct = round(net / base * 100, 2) if closed and base else None
        entry, exit_ = g["Opened Date"].min(), g["Closed Date"].max() if closed else pd.NaT
        long_rows = g[~g["Short"]]
        qty = long_rows["Buy Quantity"].map(num)
        px = long_rows["Purchase PRICE PER SHARE"].map(num)
        avg_px = round((qty * px).sum() / qty.sum(), 4) if qty.sum() > 0 else None
        acts = sorted({a for a in g["ActionRaw"] if a})
        out.append({
            "TradeID": tid, "Account": first["Account"], "Asset": first["Asset"],
            "Style": style, "StyleFrom": src, "Underlying": first["Underlying"],
            "Symbol": first["Symbol"], "Direction": "Short" if g["Short"].any() else "Long",
            "Status": "Closed" if closed else "Open",
            "EntryDate": entry, "ExitDate": exit_,
            "HoldTradingDays": trading_days(entry, exit_) if closed else pd.NA,
            "Rows": len(g),
            "Entries": int(g["Opened Date"].dropna().dt.normalize().nunique()),
            "Exits": int(g.loc[g["Closed"], "Closed Date"].dropna().dt.normalize().nunique()),
            "AvgEntryPrice": avg_px,
            "TotalInvestment": round(invest, 2), "TotalSaleValue": round(sale, 2),
            "NetProfit": round(net, 2) if closed else None,
            "NetPct": pct, "Outcome": outcome(pct, first["Asset"]) if closed else "",
            "ActionValues": ", ".join(acts),
            "Pre2026": bool(not pd.isna(entry) and entry < pd.Timestamp("2026-01-01")),
        })
        if len(acts) > 1:
            review.append(("Mixed Action values in one trade", tid, first["Account"],
                           first["Symbol"], f"{', '.join(acts)}; using '{style}'"))
        if g["Short"].any() and acts and STYLE_MAP.get(acts[0].lower()) != "Hedge":
            review.append(("Sold to open but Action is not Hedge", tid, first["Account"],
                           first["Symbol"], f"Action '{acts[0]}'; graded as Hedge"))
        if pd.isna(first["Opened Date"]):
            review.append(("No Opened Date", tid, first["Account"], first["Symbol"],
                           f"Excel row {int(first['ExcelRow'])}; can't be matched or timed"))
    trades = pd.DataFrame(out)
    trades["HoldTradingDays"] = trades["HoldTradingDays"].astype("Int64")

    # one decision = same symbol, same entry day, same style, across accounts
    key = (trades["Symbol"] + "|" + trades["EntryDate"].dt.strftime("%Y-%m-%d").fillna(trades["TradeID"])
           + "|" + trades["Style"])
    trades["DecisionID"] = [f"D{i + 1:04d}" for i in pd.factorize(key)[0]]
    trades["Accounts"] = trades.groupby("DecisionID")["Account"].transform("count")

    for t in trades.itertuples(index=False):
        if t.Style == "Untagged":
            review.append(("No style", t.TradeID, t.Account, t.Symbol,
                           "Add Swing/Long/Hedge/Day in the Action column or aar_tags.csv"))
        elif t.Style not in ("Swing", "Long term", "Hedge", "Day"):
            review.append((f"Action '{t.Style}' not mapped to a style", t.TradeID, t.Account,
                           t.Symbol, "Reported as its own group for now"))
        if t.Style == "Swing" and t.Status == "Closed" and (t.HoldTradingDays or 0) > LONG_SWING_DAYS:
            review.append(("Swing held a long time", t.TradeID, t.Account, t.Symbol,
                           f"{t.HoldTradingDays} trading days; Long term?"))
    review = pd.DataFrame(review, columns=["Issue", "TradeID", "Account", "Symbol", "Detail"])
    return trades, review


def decision_view(trades):
    """Closed trades rolled up to one row per decision (combined across accounts)."""
    c = trades[trades["Status"] == "Closed"]
    g = c.groupby("DecisionID")
    d = g.agg(Asset=("Asset", "first"), Style=("Style", "first"), Symbol=("Symbol", "first"),
              EntryDate=("EntryDate", "first"), Accounts=("Account", "count"),
              Invest=("TotalInvestment", "sum"), Sale=("TotalSaleValue", "sum"),
              Net=("NetProfit", "sum")).reset_index()
    carry = [c for c in ("Source", "FlagStatus", "EpisodeID", "Score", "Combo", "DaysAfterFlag",
                         "DaysSinceFirstFlag", "EntryType", "ActualRiskPct", "SPYAbove50",
                         "QQQAbove50", "HoldTradingDays", "DaysToExpiry", "DTEBucket",
                         "UnderMovePct", "MFEPct", "MAEPct", "StopHit", "HighestTarget", "UnderR")
             if c in trades.columns]
    if carry:
        d = d.merge(g[carry].first().reset_index(), on="DecisionID", how="left")
    base = np.where(d["Invest"] > 0, d["Invest"], d["Sale"])
    d["NetPct"] = np.round(np.where(base != 0, d["Net"] / np.where(base == 0, 1, base) * 100, np.nan), 2)
    d["Outcome"] = [outcome(p, a) for p, a in zip(d["NetPct"], d["Asset"])]
    if "ActualRiskPct" in d:
        risk = pd.to_numeric(d["ActualRiskPct"], errors="coerce")
        d["R"] = (d["NetPct"] / risk).where(risk > 0).round(2)
    if "UnderR" in d:   # options: R measured on the underlying
        d["R"] = d["R"].where(d["Asset"] == "Stock", pd.to_numeric(d["UnderR"], errors="coerce"))
    return d


# ------------------------------------------------------------ scanner match
def combo_name(r):
    parts = [n for n, col in (("Squeeze", "BBSqueeze"), ("ATR", "ATRContract"), ("Volume", "VolDryup"))
             if num(r.get(col)) > 0]
    return " + ".join(parts) if parts else "None"


def load_flags():
    """Scan log rows with episode IDs. Empty frame if there is no log yet."""
    if not SCAN_LOG.exists():
        return pd.DataFrame()
    f = pd.read_csv(SCAN_LOG)
    if f.empty:
        return f
    f["ScanDate"] = pd.to_datetime(f["ScanDate"], errors="coerce")
    f = f.dropna(subset=["ScanDate"])
    f["Symbol"] = f["Symbol"].astype(str).str.strip().str.upper()
    f["Combo"] = f.apply(combo_name, axis=1)
    f["VolRatio"] = (f["Vol10"] / f["Vol50"]).round(3)
    f = f.sort_values(["Symbol", "ScanDate"]).reset_index(drop=True)
    ep, n, prev_sym, prev_date = [], 0, None, None
    for sym, d in zip(f["Symbol"], f["ScanDate"]):
        if sym != prev_sym or trading_days(prev_date, d) > EPISODE_GAP:
            n += 1
        ep.append(f"E{n:04d}")
        prev_sym, prev_date = sym, d
    f["EpisodeID"] = ep
    f["EpisodeStart"] = f.groupby("EpisodeID")["ScanDate"].transform("min")
    return f


FLAG_FIELDS = {  # scan log column -> trade column
    "ScanDate": "FlagDate", "Status": "FlagStatus", "EpisodeID": "EpisodeID",
    "EpisodeStart": "FirstFlagDate", "Score": "Score", "Checks": "Checks", "Combo": "Combo",
    "BBWidthPctile": "BBWidthPctile", "ATRRatio": "ATRRatio", "VolRatio": "VolRatio",
    "PctBelowPivot": "PctBelowPivot", "Close": "FlagClose", "Pivot": "Pivot",
    "BaseLow": "BaseLow", "Entry": "ScanEntry", "Stop": "ScanStop", "RiskPct": "ScanRiskPct",
    "Target": "ScanTarget", "Target2R": "Target2R", "Target3R": "Target3R",
    "TargetMeasured": "TargetMeasured", "Resistance": "Resistance", "TargetATR": "TargetATR",
    "RewardRisk": "ScanRewardRisk", "RROk": "RROk", "SPYAbove50": "SPYAbove50",
    "QQQAbove50": "QQQAbove50", "Sector": "Sector", "List": "ScanList",
}


def match_trades(trades, flags, review_rows):
    """Adds Source and the matched flag's details to each trade."""
    t = trades.copy()
    for col in ["Source", "DaysAfterFlag", "DaysSinceFirstFlag", *FLAG_FIELDS.values(),
                "EntryType", "ActualRiskPct", "R"]:
        t[col] = pd.NA
    first_scan = flags["ScanDate"].min() if len(flags) else None
    by_sym = {s: g for s, g in flags.groupby("Symbol")} if len(flags) else {}
    lag_rows = []
    for i, tr in t.iterrows():
        entry = tr["EntryDate"]
        if pd.isna(entry):
            t.at[i, "Source"] = "Unknown (no entry date)"
            continue
        if first_scan is None or trading_days(first_scan, entry) < 1:
            t.at[i, "Source"] = "Before scan log"
            continue
        g = by_sym.get(tr["Underlying"])
        hit = None
        if g is not None:
            prior = g[g["ScanDate"] < entry.normalize()]
            if len(prior):
                last = prior.iloc[-1]
                lag = trading_days(last["ScanDate"], entry)
                if 1 <= lag <= MATCH_DAYS:
                    hit = last
                elif lag <= MATCH_DAYS * 2:   # just outside the window: helps tune it
                    review_rows.append(("Flagged just outside the match window", tr["TradeID"],
                                        tr["Account"], tr["Symbol"],
                                        f"last flag {last['ScanDate']:%Y-%m-%d} ({last['Status']}), "
                                        f"{lag} trading days before entry"))
        if hit is None:
            t.at[i, "Source"] = "Discretionary"
            continue
        t.at[i, "Source"] = "Scanner"
        for src, dst in FLAG_FIELDS.items():
            t.at[i, dst] = hit[src]
        t.at[i, "DaysAfterFlag"] = trading_days(hit["ScanDate"], entry)
        t.at[i, "DaysSinceFirstFlag"] = trading_days(hit["EpisodeStart"], entry)
        lag_rows.append(t.at[i, "DaysAfterFlag"])
        # stock-only measures (option premiums can't be compared to the stock's pivot)
        px = tr["AvgEntryPrice"]
        if tr["Asset"] == "Stock" and not blank(px) and px:
            if not blank(hit.get("Pivot")):
                t.at[i, "EntryType"] = "Pre-breakout" if px <= hit["Pivot"] else "Post-breakout"
            stop = hit.get("Stop")
            if not blank(stop) and px > stop:
                risk = (px - stop) / px * 100
                t.at[i, "ActualRiskPct"] = round(risk, 2)
                if tr["Status"] == "Closed" and not blank(tr["NetPct"]):
                    t.at[i, "R"] = round(tr["NetPct"] / risk, 2)
    for c in ("DaysAfterFlag", "DaysSinceFirstFlag", "Score", "Checks"):
        t[c] = pd.to_numeric(t[c], errors="coerce").astype("Int64")
    for c in ("FlagDate", "FirstFlagDate"):
        t[c] = pd.to_datetime(t[c], errors="coerce")
    return t


def flag_table(flags, trades, as_of=None):
    """One row per setup episode: what the scanner saw and whether it was taken.
    'Window open' = the match window after the last flag hasn't closed yet."""
    if flags.empty:
        return pd.DataFrame()
    today = pd.Timestamp(as_of or pd.Timestamp.today()).normalize()
    first = flags.sort_values("ScanDate").groupby("EpisodeID").first()
    g = flags.groupby("EpisodeID")
    ep = pd.DataFrame({
        "Symbol": first["Symbol"], "FirstFlagDate": first["ScanDate"],
        "LastFlagDate": g["ScanDate"].max(), "Flags": g.size(),
        "BestStatus": g["Status"].agg(lambda s: max(s, key=lambda x: STATUS_RANK.get(x, 0))),
        "Statuses": g["Status"].agg(lambda s: ", ".join(dict.fromkeys(s))),
        "MaxScore": g["Score"].max(),
        "FirstStatus": first["Status"], "FirstScore": first["Score"], "Combo": first["Combo"],
        "BBWidthPctile": first["BBWidthPctile"], "ATRRatio": first["ATRRatio"],
        "VolRatio": first["VolRatio"], "PctBelowPivot": first["PctBelowPivot"],
        "Pivot": first["Pivot"], "ScanEntry": first.get("Entry"), "ScanStop": first.get("Stop"),
        "ScanTarget": first.get("Target"), "ScanRewardRisk": first.get("RewardRisk"),
        **{c: first.get(c) for c, _ in TARGET_COLS},
        "SPYAbove50": first["SPYAbove50"], "QQQAbove50": first["QQQAbove50"],
        "Sector": first["Sector"],
    }).reset_index()
    m = trades[trades["Source"] == "Scanner"]
    taken = m.groupby("EpisodeID").agg(
        TradeIDs=("TradeID", lambda s: ", ".join(s)),
        DecisionIDs=("DecisionID", lambda s: ", ".join(dict.fromkeys(s))),
        Styles=("Style", lambda s: ", ".join(dict.fromkeys(s))),
        Assets=("Asset", lambda s: ", ".join(dict.fromkeys(s))))
    ep = ep.merge(taken, left_on="EpisodeID", right_index=True, how="left")
    open_window = ep["LastFlagDate"].map(lambda d: trading_days(d, today) <= MATCH_DAYS)
    ep["Taken"] = np.where(ep["TradeIDs"].notna(), "Taken",
                           np.where(open_window, "Window open", "Passed"))
    return ep.sort_values(["FirstFlagDate", "MaxScore"], ascending=[False, False])


# ------------------------------------------------------------------- prices
def last_completed_session(now=None):
    """Most recent trading day whose daily bar is final (after 4:15 PM ET)."""
    now = pd.Timestamp(now) if now is not None else pd.Timestamp.now(tz="America/New_York")
    if now.tzinfo is None:
        now = now.tz_localize("America/New_York")
    d = np.datetime64(now.date(), "D")
    if np.is_busday(d, holidays=NYSE_HOLIDAYS):
        if now.time() >= dt.time(16, 15):
            return pd.Timestamp(d)
        return pd.Timestamp(np.busday_offset(d, -1, holidays=NYSE_HOLIDAYS))
    return pd.Timestamp(np.busday_offset(d, 0, roll="backward", holidays=NYSE_HOLIDAYS))


def add_trading_days(date, n):
    d = np.busday_offset(np.datetime64(pd.Timestamp(date).date(), "D"), n, roll="forward",
                         holidays=NYSE_HOLIDAYS)
    return pd.Timestamp(d)


class PriceBook:
    """Daily bars per symbol, cached in data/aar_bars/<SYMBOL>.csv. A symbol is
    only fetched when the cache doesn't reach the date a calculation needs, so
    finished trades and setups never cost another request."""

    def __init__(self, source=None, enabled=True, now=None):
        self.enabled = enabled
        self._source = source
        self._mem = {}
        self.fetched, self.failed, self.stopped = [], {}, None
        self.through = last_completed_session(now)
        self._t0 = time.monotonic()

    def source(self):
        if self._source is None:
            from data import get_source   # read-only price history, same as the scanner
            self._source = get_source()
        return self._source

    def _path(self, sym):
        return BARS_DIR / f"{sym.replace('/', '_').replace('.', '_')}.csv"

    def _cached(self, sym):
        if sym in self._mem:
            return self._mem[sym]
        p = self._path(sym)
        df = None
        if p.exists():
            df = pd.read_csv(p, index_col=0, parse_dates=True)
        self._mem[sym] = df
        return df

    def bars(self, sym, need_through):
        """Bars for sym, fetching only if the cache ends before need_through."""
        need = min(pd.Timestamp(need_through).normalize(), self.through)
        df = self._cached(sym)
        if df is not None and len(df) and df.index.max() >= need:
            return df
        if not self.enabled or self.stopped or sym in self.failed:
            return df
        try:
            new = self.source().history(sym)
        except Exception as e:                       # login expired stops all fetching
            if type(e).__name__ == "DataError":
                self.stopped = str(e)
            else:
                self.failed[sym] = str(e)[:120]
            return df
        if new is None or new.empty:
            self.failed[sym] = "no price history returned"
            return df
        new = new[["open", "high", "low", "close", "volume"]].copy()
        new.index = pd.to_datetime(new.index).normalize()
        merged = new if df is None else pd.concat([df[~df.index.isin(new.index)], new]).sort_index()
        BARS_DIR.mkdir(parents=True, exist_ok=True)
        merged.to_csv(self._path(sym))
        self._mem[sym] = merged
        self.fetched.append(sym)
        n = len(self.fetched)
        if n == 1 or n % 10 == 0:
            secs = time.monotonic() - self._t0
            print(f"    {n} symbols fetched ({secs:.0f}s, latest {sym})", flush=True)
        return merged


def window(df, start, end):
    if df is None or df.empty:
        return df
    return df[(df.index >= pd.Timestamp(start).normalize()) & (df.index <= pd.Timestamp(end).normalize())]


def dte_bucket(days):
    if days is None or pd.isna(days):
        return pd.NA
    for limit, label in DTE_BUCKETS:
        if days <= limit:
            return label
    return pd.NA


def add_option_dte(trades):
    t = trades.copy()
    exp = t["Symbol"].map(option_expiry)
    dte = [(e - d).days if e is not None and not pd.isna(d) and a == "Option" else pd.NA
           for e, d, a in zip(exp, t["EntryDate"], t["Asset"])]
    t["DaysToExpiry"] = pd.array(dte, dtype="Int64")
    t["DTEBucket"] = [dte_bucket(x) for x in dte]
    return t


def trade_paths(trades, book):
    """What the underlying did while each Swing/Day trade was open (daily bars,
    entry day through exit day). Stocks are measured from your average entry
    price, options from the underlying's close on the entry day."""
    t = trades.copy()
    new_cols = ["UnderEntry", "UnderExit", "UnderMovePct", "MFEPct", "MAEPct", "PathNote",
                "StopHit", "HighestTarget", "UnderR"] + [f"Hit{lbl.replace(' ', '')}" for _, lbl in TARGET_COLS]
    for c in new_cols:
        t[c] = pd.NA
    for i, tr in t.iterrows():
        if tr["Style"] not in ("Swing", "Day") or tr["Status"] != "Closed" or pd.isna(tr["EntryDate"]):
            continue
        df = book.bars(tr["Underlying"], tr["ExitDate"])
        w = window(df, tr["EntryDate"], tr["ExitDate"])
        if w is None or w.empty or w.index.min() > pd.Timestamp(tr["EntryDate"]).normalize():
            t.at[i, "PathNote"] = "no price history for these dates"
            continue
        ref_close = float(w["close"].iloc[0])
        ref = float(tr["AvgEntryPrice"]) if tr["Asset"] == "Stock" and not blank(tr["AvgEntryPrice"]) \
            else ref_close
        t.at[i, "UnderEntry"] = round(ref_close, 2)
        t.at[i, "UnderExit"] = round(float(w["close"].iloc[-1]), 2)
        t.at[i, "UnderMovePct"] = round((w["close"].iloc[-1] / ref_close - 1) * 100, 2)
        t.at[i, "MFEPct"] = round((w["high"].max() / ref - 1) * 100, 2)
        t.at[i, "MAEPct"] = round((w["low"].min() / ref - 1) * 100, 2)
        t.at[i, "PathNote"] = "daily bars, entry day to exit day"
        if tr.get("Source") != "Scanner":
            continue
        stop = tr.get("ScanStop")
        if not blank(stop):
            t.at[i, "StopHit"] = bool(w["low"].min() <= float(stop))
            if tr["Asset"] == "Option" and ref_close > float(stop):
                t.at[i, "UnderR"] = round((w["close"].iloc[-1] - ref_close) / (ref_close - float(stop)), 2)
            if tr["Asset"] == "Option" and not blank(tr.get("Pivot")):
                t.at[i, "EntryType"] = "Pre-breakout" if ref_close <= float(tr["Pivot"]) else "Post-breakout"
        best, best_lvl = None, -1
        for col, lbl in TARGET_COLS:
            lvl = tr.get(col)
            if blank(lvl):
                continue
            hit = bool(w["high"].max() >= float(lvl))
            t.at[i, f"Hit{lbl.replace(' ', '')}"] = hit
            if hit and float(lvl) > best_lvl:
                best, best_lvl = lbl, float(lvl)
        t.at[i, "HighestTarget"] = best or ("none" if not blank(stop) else pd.NA)
    return t


OPEN_COLS = ["LastDate", "LastClose", "DaysHeld", "OpenMovePct", "OpenMFEPct", "OpenMAEPct",
             "OpenR", "TargetR", "ToStopPct", "ToTargetPct", "OpenStopHit", "OpenHighestTarget"]


def open_positions(trades, book):
    """Where each open swing trade stands as of the last completed session
    (closing prices, not the tracker; nothing here is P&L). Only trades since
    the scan log began. Stocks are measured from your average entry price,
    options from the underlying's close on the entry day. R uses the
    scanner's stop when the trade matched a flag."""
    t = trades.copy()
    for c in OPEN_COLS:
        t[c] = pd.NA
    for i, tr in t.iterrows():
        if (tr["Status"] != "Open" or tr["Style"] != "Swing" or pd.isna(tr["EntryDate"])
                or tr.get("Source") not in ("Scanner", "Discretionary")):
            continue
        df = book.bars(tr["Underlying"], book.through)
        w = window(df, tr["EntryDate"], book.through)
        if w is None or w.empty or w.index.min() > pd.Timestamp(tr["EntryDate"]).normalize():
            continue
        last = float(w["close"].iloc[-1])
        ref = float(tr["AvgEntryPrice"]) if tr["Asset"] == "Stock" and not blank(tr["AvgEntryPrice"]) \
            else float(w["close"].iloc[0])
        t.at[i, "LastDate"] = w.index[-1]
        t.at[i, "LastClose"] = round(last, 2)
        t.at[i, "DaysHeld"] = trading_days(tr["EntryDate"], w.index[-1])
        t.at[i, "OpenMovePct"] = round((last / ref - 1) * 100, 2)
        t.at[i, "OpenMFEPct"] = round((w["high"].max() / ref - 1) * 100, 2)
        t.at[i, "OpenMAEPct"] = round((w["low"].min() / ref - 1) * 100, 2)
        stop, target = tr.get("ScanStop"), tr.get("ScanTarget")
        if not blank(stop) and ref > float(stop):
            risk = ref - float(stop)
            t.at[i, "OpenR"] = round((last - ref) / risk, 2)
            t.at[i, "ToStopPct"] = round((float(stop) / last - 1) * 100, 2)
            t.at[i, "OpenStopHit"] = bool(w["low"].min() <= float(stop))
            if not blank(target):
                t.at[i, "TargetR"] = round((float(target) - ref) / risk, 2)
                t.at[i, "ToTargetPct"] = round((float(target) / last - 1) * 100, 2)
        best, best_lvl = None, -1
        for col, lbl in TARGET_COLS:
            lvl = tr.get(col)
            if not blank(lvl) and w["high"].max() >= float(lvl) and float(lvl) > best_lvl:
                best, best_lvl = lbl, float(lvl)
        if not blank(stop):
            t.at[i, "OpenHighestTarget"] = best or "none yet"
    t["LastDate"] = pd.to_datetime(t["LastDate"], errors="coerce")
    t["DaysHeld"] = pd.to_numeric(t["DaysHeld"], errors="coerce").astype("Int64")
    return t


def flag_outcomes(ep, book):
    """How each setup episode played out on its own, whatever you did with it.
    Levels come from the episode's first flag. Triggered = price traded at or
    above Entry within TRIGGER_DAYS sessions after the flag. Then the first of
    Stop or Target (a day touching both counts as Stop), the highest target
    reached, and the close OUTCOME_DAYS sessions after the trigger."""
    if ep is None or ep.empty:
        return ep
    e = ep.copy()
    cols = ["Triggered", "TriggerDate", "Result", "DaysToResult", "HighestTarget",
            "MFEPct", "MAEPct", "Ret10Pct", "Ret20Pct", "R20", "OutcomeStatus"]
    for c in cols:
        e[c] = pd.NA
    for i, r in e.iterrows():
        entry = r["ScanEntry"] if not blank(r.get("ScanEntry")) else r.get("Pivot")
        if blank(entry):
            e.at[i, "OutcomeStatus"] = "no entry level"
            continue
        entry, stop, target = float(entry), r.get("ScanStop"), r.get("ScanTarget")
        flag_day = pd.Timestamp(r["FirstFlagDate"])
        need = add_trading_days(flag_day, TRIGGER_DAYS + OUTCOME_DAYS + 1)
        df = book.bars(r["Symbol"], need)
        after = window(df, add_trading_days(flag_day, 1), need)
        if after is None or after.empty:
            e.at[i, "OutcomeStatus"] = "pending (no bars after flag yet)" if df is not None \
                else "no price history"
            continue
        trig = after.iloc[:TRIGGER_DAYS]
        hit = trig.index[trig["high"] >= entry]
        if len(hit) == 0:
            e.at[i, "Triggered"] = False
            done = len(trig) >= TRIGGER_DAYS
            e.at[i, "Result"] = "Not triggered" if done else pd.NA
            e.at[i, "OutcomeStatus"] = "complete" if done else "pending"
            continue
        tday = hit[0]
        e.at[i, "Triggered"], e.at[i, "TriggerDate"] = True, tday
        path = after[after.index >= tday].iloc[:OUTCOME_DAYS + 1]   # trigger day + N sessions
        result, n = None, None
        for k, (d, b) in enumerate(path.iterrows()):
            stop_hit = not blank(stop) and b["low"] <= float(stop)
            tgt_hit = not blank(target) and b["high"] >= float(target)
            if stop_hit:
                result, n = "Stop first", k
                break
            if tgt_hit:
                result, n = "Target first", k
                break
        complete = len(path) >= OUTCOME_DAYS + 1
        e.at[i, "Result"] = result or ("Neither (open)" if not complete else f"Neither in {OUTCOME_DAYS}d")
        e.at[i, "DaysToResult"] = n
        # targets only count if reached before the stop day (the stop day's high
        # may have come after the stop, so it is ignored: conservative)
        upto = path if result != "Stop first" else path.iloc[:n]
        best, best_lvl = "none", -1
        for col, lbl in TARGET_COLS:
            lvl = r.get(col)
            if blank(lvl) or upto.empty:
                continue
            if upto["high"].max() >= float(lvl) and float(lvl) > best_lvl:
                best, best_lvl = lbl, float(lvl)
        e.at[i, "HighestTarget"] = best
        e.at[i, "MFEPct"] = round((path["high"].max() / entry - 1) * 100, 2)
        e.at[i, "MAEPct"] = round((path["low"].min() / entry - 1) * 100, 2)
        if len(path) > 10:
            e.at[i, "Ret10Pct"] = round((path["close"].iloc[10] / entry - 1) * 100, 2)
        if complete:
            ret20 = (path["close"].iloc[OUTCOME_DAYS] / entry - 1) * 100
            e.at[i, "Ret20Pct"] = round(ret20, 2)
            if not blank(stop) and entry > float(stop):
                e.at[i, "R20"] = round(ret20 / ((entry - float(stop)) / entry * 100), 2)
        e.at[i, "OutcomeStatus"] = "complete" if complete else ("resolved, 20d return pending"
                                                                if result else "pending")
    e["TriggerDate"] = pd.to_datetime(e["TriggerDate"], errors="coerce")
    return e


# ------------------------------------------------------------------ summary
def summarize(df, by, report):
    """Win/Scratch/Loss table for closed decisions, one row per group."""
    rows = []
    if df.empty:
        return pd.DataFrame()
    for key, g in df.groupby(by, sort=True, dropna=False):
        key = key if isinstance(key, tuple) else (key,)
        p = g["NetPct"].astype(float)
        wins, losses = p[g["Outcome"] == "Win"], p[g["Outcome"] == "Loss"]
        n = len(g)
        r = g["R"].astype(float).dropna() if "R" in g else pd.Series(dtype=float)
        rows.append({"Report": report,
                     "Group": " | ".join("(blank)" if pd.isna(k) else str(k) for k in key), "N": n,
                     "Win%": round((g["Outcome"] == "Win").mean() * 100, 1),
                     "Scratch%": round((g["Outcome"] == "Scratch").mean() * 100, 1),
                     "Loss%": round((g["Outcome"] == "Loss").mean() * 100, 1),
                     "AvgPct": round(p.mean(), 2), "MedianPct": round(p.median(), 2),
                     "AvgWinPct": round(wins.mean(), 2) if len(wins) else None,
                     "AvgLossPct": round(losses.mean(), 2) if len(losses) else None,
                     "AvgR": round(r.mean(), 2) if len(r) else None,
                     "Note": "too few to read" if n < 20 else ""})
    return pd.DataFrame(rows)


def summarize_flags(df, by, report):
    """How setups played out on their own (from Flags), one row per group."""
    if df is None or df.empty or "Triggered" not in df:
        return pd.DataFrame()
    done = df[df["OutcomeStatus"].isin(["complete", "resolved, 20d return pending"])
              | (df["Result"] == "Not triggered")]
    rows = []
    for key, g in done.groupby(by, sort=True, dropna=False):
        key = key if isinstance(key, tuple) else (key,)
        trig = g[g["Triggered"] == True]  # noqa: E712
        res = trig[trig["Result"].isin(["Target first", "Stop first"])]
        r20 = pd.to_numeric(trig["Ret20Pct"], errors="coerce").dropna()
        rr = pd.to_numeric(trig["R20"], errors="coerce").dropna()
        rows.append({"Report": report,
                     "Group": " | ".join("(blank)" if pd.isna(k) else str(k) for k in key),
                     "N": len(g), "Triggered%": round(len(trig) / len(g) * 100, 1),
                     "TargetFirst%": round((res["Result"] == "Target first").mean() * 100, 1)
                     if len(res) else None,
                     "AvgRet20Pct": round(r20.mean(), 2) if len(r20) else None,
                     "MedianRet20Pct": round(r20.median(), 2) if len(r20) else None,
                     "AvgR20": round(rr.mean(), 2) if len(rr) else None,
                     "Note": "too few to read" if len(g) < 20 else ""})
    return pd.DataFrame(rows)


def build_summary(decisions, flags_ep, trades):
    d = decisions
    swing = d[d["Style"] == "Swing"]
    live = swing[swing["Source"].isin(["Scanner", "Discretionary"])]
    scan = swing[swing["Source"] == "Scanner"]
    parts = [
        summarize(d, ["Style", "Asset"], "All closed decisions by style"),
        summarize(live, ["Asset", "Source"], "Swing: scanner vs discretionary (since scan log began)"),
        summarize(scan, ["Asset", "FlagStatus"], "Swing scanner trades: flag status"),
        summarize(scan, ["Asset", "Combo"], "Swing scanner trades: check combination"),
        summarize(scan, ["Asset", "Score"], "Swing scanner trades: score"),
        summarize(scan, ["Asset", "SPYAbove50", "QQQAbove50"],
                  "Swing scanner trades: SPY | QQQ above 50-day"),
        summarize(scan[scan["Asset"] == "Stock"], ["EntryType"],
                  "Swing scanner stock trades: entry vs pivot"),
        summarize(scan, ["Asset", "DaysAfterFlag"], "Swing scanner trades: days after flag"),
        summarize(swing[swing["Asset"] == "Option"], ["DTEBucket"],
                  "Swing options: days to expiry at entry"),
        summarize(swing[swing["Asset"] == "Option"].assign(
            UnderDir=lambda x: pd.cut(pd.to_numeric(x["UnderMovePct"], errors="coerce"),
                                      [-1e9, -2, 2, 1e9], labels=["Down >2%", "Flat", "Up >2%"])),
                  ["UnderDir"], "Swing options: underlying move while held"),
    ]
    if len(flags_ep) and "Triggered" in flags_ep:
        parts += [
            summarize_flags(flags_ep, ["BestStatus"], "Setups on their own: status (tests 2-of-3)"),
            summarize_flags(flags_ep, ["Combo"], "Setups on their own: check combination"),
            summarize_flags(flags_ep, ["FirstScore"], "Setups on their own: score"),
            summarize_flags(flags_ep, ["SPYAbove50", "QQQAbove50"],
                            "Setups on their own: SPY | QQQ above 50-day"),
            summarize_flags(flags_ep, ["Taken"], "Setups on their own: taken vs passed"),
        ]
    out = pd.concat([p for p in parts if len(p)], ignore_index=True)
    if len(flags_ep):
        cnt = (flags_ep.groupby(["BestStatus", "Taken"]).size().rename("N").reset_index())
        cnt = cnt.assign(Report="Setup episodes: taken vs passed (counts)",
                         Group=cnt["BestStatus"] + " | " + cnt["Taken"])[["Report", "Group", "N"]]
        out = pd.concat([out, cnt], ignore_index=True)
    return out


# ------------------------------------------------------------------- output
def write_excel(sheets):
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=OUTPUT.parent)
    os.close(fd)
    with pd.ExcelWriter(tmp, engine="openpyxl") as xw:
        for name, df in sheets.items():
            df = df.copy()
            for c in df.columns:  # dates without the time part
                if pd.api.types.is_datetime64_any_dtype(df[c]):
                    df[c] = df[c].dt.date
            df.to_excel(xw, sheet_name=name, index=False)
            ws = xw.sheets[name]
            for i, col in enumerate(df.columns, 1):
                width = max([len(str(col))] + [len(str(v)) for v in df[col].head(300)]) + 2
                ws.column_dimensions[get_column_letter(i)].width = min(width, 40)
            ws.freeze_panes = "A2"
            if len(df):
                ref = f"A1:{get_column_letter(len(df.columns))}{len(df) + 1}"
                tbl = Table(displayName=name.replace(" ", ""), ref=ref)
                tbl.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
                ws.add_table(tbl)
    try:
        os.replace(tmp, OUTPUT)
    except PermissionError:
        os.remove(tmp)
        sys.exit(f"{OUTPUT.name} is open in Excel. Close it and run again.")


def scan_log_info():
    if not SCAN_LOG.exists():
        return "not found"
    s = pd.read_csv(SCAN_LOG, usecols=["ScanDate", "Status"])
    return (f"{len(s)} rows, {s['ScanDate'].min()} to {s['ScanDate'].max()} "
            f"({', '.join(f'{v} {k}' for k, v in s['Status'].value_counts().items())})")


class Tee:
    """Writes to the console and appends to a log file, so a run shows its
    progress on screen and still leaves a record."""

    def __init__(self, stream, path):
        self.stream, self.f = stream, open(path, "a", encoding="utf-8")

    def write(self, s):
        self.stream.write(s)
        self.f.write(s)
        return len(s)

    def flush(self):
        self.stream.flush()
        self.f.flush()


def main():
    ap = argparse.ArgumentParser(description="After Action Review: grade trades from the tracker")
    ap.add_argument("--tracker", help="tracker file (default: newest in TRACKER_DIR)")
    ap.add_argument("--no-prices", action="store_true",
                    help="skip price history (no Schwab calls; uses cached bars only)")
    ap.add_argument("--source", help="price source: schwab or yahoo (default: DATA_SOURCE)")
    ap.add_argument("--log", help="also append everything printed to this file (run_aar.bat uses aar_log.txt)")
    args = ap.parse_args()
    if args.log:
        sys.stdout = Tee(sys.stdout, HERE / args.log)
        sys.stderr = Tee(sys.stderr, HERE / args.log)
        print(f"==== {dt.datetime.now():%a %m/%d/%Y %H:%M}")
    started = time.monotonic()

    path = Path(args.tracker) if args.tracker else find_latest_tracker()
    if not path.is_absolute():
        path = TRACKER_DIR / path
    trk, name = read_tracker(path)
    rows = chain_trades(prepare_rows(trk, short_open_fills()))
    trades, review = build_trades(rows, load_tags())
    flags = load_flags()
    extra = []
    trades = match_trades(trades, flags, extra)
    trades = add_option_dte(trades)
    flags_ep = flag_table(flags, trades)

    source = None
    if args.source and not args.no_prices:
        from data import get_source
        source = get_source(args.source)
    book = PriceBook(source=source, enabled=not args.no_prices)
    print(f"After Action Review  ({name})")
    if not args.no_prices:
        print("  Price history: checking cache, fetching only what's missing "
              "(a few minutes when many setups are new)...", flush=True)
    trades = trade_paths(trades, book)
    trades = open_positions(trades, book)
    flags_ep = flag_outcomes(flags_ep, book)
    for sym, err in book.failed.items():
        extra.append(("Price history unavailable", "", "", sym, err))
    if extra:
        review = pd.concat([review, pd.DataFrame(extra, columns=review.columns)], ignore_index=True)
    decisions = decision_view(trades)
    summary = build_summary(decisions, flags_ep, trades)

    row_cols = ["TradeID", "ExcelRow", "Account", "Opened Date", "Stock Or Option", "Action",
                "SYMBOL", "Buy Quantity", "Purchase PRICE PER SHARE", "TOTAL INVESTMENT",
                "Closed Date", "Sell Quantity", "Total Sale Value", "% CHANGE", "Net Profit",
                "Open/Closed", "Sync ID", "Short", "RowStyle"]
    price_note = ("skipped (--no-prices), cached bars only" if args.no_prices else
                  f"{len(book.fetched)} symbols fetched, {len(book.failed)} unavailable"
                  + (f"; STOPPED: {book.stopped}" if book.stopped else ""))
    info = pd.DataFrame({"Item": ["Run", "Tracker", "Tracker rows", "Trades", "Decisions (closed)",
                                  "Scan log", "Price history", "Bars complete through",
                                  "Scratch band stock %", "Scratch band option %",
                                  "Match window (trading days)", "Episode gap",
                                  "Trigger window / outcome window", "Build step"],
                         "Value": [dt.datetime.now().strftime("%Y-%m-%d %H:%M"), name, len(rows),
                                   len(trades), len(decisions), scan_log_info(), price_note,
                                   str(book.through.date()), SCRATCH["Stock"], SCRATCH["Option"],
                                   f"1 to {MATCH_DAYS}", EPISODE_GAP,
                                   f"{TRIGGER_DAYS} / {OUTCOME_DAYS} sessions",
                                   "3: price paths, targets, option DTE, setup outcomes"]})
    write_excel({"Summary": summary, "Trades": trades, "Decisions": decisions, "Flags": flags_ep,
                 "TradeRows": rows[row_cols].sort_values(["TradeID", "Opened Date"]),
                 "Review": review, "RunInfo": info})
    try:
        import aar_report
        aar_report.build(decisions, trades, flags_ep, review, info, REPORT)
    except Exception as e:   # the dashboard never blocks the workbook
        print(f"  Dashboard not built: {type(e).__name__}: {e}")

    print(f"  {len(rows)} tracker rows -> {len(trades)} trades "
          f"({(trades['Status'] == 'Closed').sum()} closed) -> {len(decisions)} closed decisions")
    print(f"  Price history: {price_note}")
    show = ("All closed", "Swing: scanner vs", "Swing options: days to expiry",
            "Swing options: underlying move", "Setups on their own: status",
            "Setups on their own: taken")
    for rep in summary["Report"].dropna().unique() if len(summary) else []:
        if not rep.startswith(show) and not rep.startswith("Setup episodes"):
            continue
        part = summary[summary["Report"] == rep].dropna(axis=1, how="all")
        keep = [c for c in ["Group", "N", "Win%", "Scratch%", "Loss%", "AvgPct", "MedianPct",
                            "Triggered%", "TargetFirst%", "AvgRet20Pct", "AvgR20", "Note"]
                if c in part]
        print(f"\n  {rep}:")
        print(part[keep].to_string(index=False))
    src = trades["Source"].value_counts()
    print("\n  Trades by source: " + ", ".join(f"{v} {k.lower()}" for k, v in src.items()))
    if book.stopped:
        print(f"\n  ** Price history stopped: {book.stopped}")
    print(f"\n  Review items: {len(review)}"
          + (f" ({', '.join(f'{v} {k.lower()}' for k, v in review['Issue'].value_counts().items())})"
             if len(review) else ""))
    print(f"  -> {OUTPUT}")
    if REPORT.exists():
        print(f"  -> {REPORT}  (open in your browser)")
    print(f"  Done in {time.monotonic() - started:.0f}s.")


if __name__ == "__main__":
    main()
