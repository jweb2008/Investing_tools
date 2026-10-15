"""
Backtest: replay the scanner over years of daily history and see how its
setups actually played out. Gives an early read on whether the scoring
works, without waiting months for the live scan log.

  py backtest.py --download     pull ~6 years of daily bars for the universe (once,
                                ~10 to 15 minutes; refreshes files older than 7 days)
  py backtest.py                run the backtest on the cached history (a minute or two)
  py backtest.py --years 3      only test the most recent 3 years
  py backtest.py --compare      compare exit rules (current target, 3R, trailing) side by side

How a trade is simulated (same rules as the live scanner and levels.py)
  Signal   any day a stock is in a base: uptrend (close > 50 > 200 SMA) and
           within 5% below the pivot. Setups (2+ contraction checks) and near
           misses (fewer checks) are both tested so the cutoff can be judged.
  Entry    buy-stop at that day's pivot, good for the next 5 trading days.
           Fills at the pivot, or at the open if the stock gaps above it.
  Stop     tighter of the 30-day base low and 2x ATR below entry, at least 1 ATR.
  Target   lower of the measured move and the nearest overhead swing high.
  Exit     stop or target, whichever is hit first (if both on the same day,
           the stop is assumed, which is conservative); otherwise at the close
           after 20 trading days. One trade at a time per stock.
  Result   R multiple = (exit - entry) / (entry - stop), plus % return.

Known limits, read results with these in mind
  - Survivorship bias: it tests today's index members. Stocks that failed
    and were dropped from the indexes are missing, which flatters results.
  - No earnings dates in history, so earnings gaps are in the results but
    cannot be separated out.
  - No commissions or slippage beyond gap fills.
  - Stocks only. Option prices are not available historically for free.

Output
  Backtest Report.xlsx in the Investing folder (summary tables), and
  data/backtest_trades.csv (every simulated trade) for your own analysis.
"""
import argparse
import datetime as dt
import os
import sys
import time
from pathlib import Path

import numpy as np
import pandas as pd

import levels
import rs
import scanner
from data import DataError, get_source

HERE = Path(__file__).parent
HIST_DIR = HERE / "data" / "history"
UNIVERSE = HERE / "data" / "universe.txt"
TRADES_CSV = HERE / "data" / "backtest_trades.csv"
REPORT = HERE / "Backtest Report.xlsx"

DOWNLOAD_YEARS = 6
ENTRY_WINDOW = 5      # trading days the buy-stop stays active
MAX_HOLD = 20         # trading days before a time exit (current rules)

# Exit rules compared by  py backtest.py --compare
#   target    the live rules: conservative target (measured move capped by
#             resistance), 20-day time exit
#   3R        fixed target at 3x the risk, 30-day time exit
#   trail10   no target: once the trade reaches +1R the stop moves to
#             breakeven, then exit on a close below the 10-day SMA; 40-day limit
#   trail20   same with the 20-day SMA (gives trades more room); 60-day limit
EXIT_MODES = {
    "target": {"target": "levels", "max_hold": 20},
    "3R": {"target": 3.0, "max_hold": 30},
    "trail10": {"target": None, "trail_sma": 10, "max_hold": 40},
    "trail20": {"target": None, "trail_sma": 20, "max_hold": 60},
}
REFRESH_DAYS = 7


# ------------------------------------------------------------------ data
def load_symbols(path=None):
    path = path or UNIVERSE
    if not path.exists():
        sys.exit(f"No universe yet ({path.name}). Run: py build_universe.py"
                 + (" and py build_momentum_universe.py" if path != UNIVERSE else ""))
    syms = [l.split("#")[0].strip().upper() for l in path.read_text().splitlines()]
    return [s for s in syms if s]


def download(years=DOWNLOAD_YEARS, source=None, symbols=None):
    src = get_source(source)
    HIST_DIR.mkdir(parents=True, exist_ok=True)
    syms = ["SPY", "QQQ"] + (symbols if symbols is not None else load_symbols())
    days = int(years * 365.25) + 10
    todo = [s for s in syms if not _fresh(s)]
    print(f"Downloading {years} years of daily bars for {len(todo)} symbols "
          f"({len(syms) - len(todo)} already fresh) with {src.name} data...")
    t0, fails = time.time(), []
    for i, s in enumerate(todo, 1):
        try:
            df = src.history(s, days=days)
            if len(df):
                df.to_csv(HIST_DIR / f"{s.replace('/', '_')}.csv")
            else:
                fails.append(s)
        except DataError as e:
            sys.exit(f"DOWNLOAD STOPPED: {e}")
        except Exception:
            fails.append(s)
        if i % 100 == 0:
            print(f"  {i}/{len(todo)}  ({time.time() - t0:.0f}s)")
    print(f"Done in {time.time() - t0:.0f}s. {len(fails)} failed"
          + (f": {', '.join(fails[:15])}" if fails else "."))


def _fresh(sym):
    p = HIST_DIR / f"{sym.replace('/', '_')}.csv"
    return p.exists() and (time.time() - p.stat().st_mtime) / 86400 < REFRESH_DAYS


def load_history(sym):
    p = HIST_DIR / f"{sym.replace('/', '_')}.csv"
    if not p.exists():
        return None
    df = pd.read_csv(p, index_col=0, parse_dates=True)
    return df[["open", "high", "low", "close", "volume"]].astype(float).sort_index()


# ------------------------------------------------------------ indicators
def indicator_frame(df: pd.DataFrame) -> pd.DataFrame:
    """scanner.indicators() for every bar at once (same formulas, vectorized)."""
    c, h, l, v = df["close"], df["high"], df["low"], df["volume"]
    out = pd.DataFrame(index=df.index)
    out["close"] = c
    out["sma50"] = c.rolling(50).mean()
    out["sma200"] = c.rolling(200).mean()
    out["pivot"] = h.shift(1).rolling(scanner.PIVOT_DAYS).max()
    out["base_low"] = l.shift(1).rolling(scanner.PIVOT_DAYS).min()
    out["pct_below_pivot"] = (out["pivot"] - c) / out["pivot"] * 100

    mid = c.rolling(scanner.BB_LEN).mean()
    sd = c.rolling(scanner.BB_LEN).std(ddof=0)
    width = (2 * scanner.BB_STD * sd) / mid
    out["bb_width_pctile"] = width.rolling(scanner.BB_LOOKBACK).rank(method="max", pct=True) * 100

    prev = c.shift(1)
    tr = pd.concat([h - l, (h - prev).abs(), (l - prev).abs()], axis=1).max(axis=1)
    tr_pct = tr / c * 100
    atr_recent = tr_pct.rolling(scanner.ATR_RECENT).mean()
    atr_prior = tr_pct.shift(scanner.ATR_RECENT).rolling(scanner.ATR_PRIOR).mean()
    out["atr_ratio"] = atr_recent / atr_prior
    out["atr14"] = tr.ewm(alpha=1 / 14, adjust=False).mean()
    out["vol_10"] = v.rolling(scanner.VOL_SHORT).mean()
    out["vol_50"] = v.rolling(scanner.VOL_LONG).mean()
    out["dollar_vol"] = (c * v).rolling(scanner.VOL_LONG).mean()
    out["bars"] = np.arange(1, len(df) + 1)

    out["uptrend"] = (c > out["sma50"]) & (out["sma50"] > out["sma200"])
    out["near_pivot"] = out["pct_below_pivot"].between(0, scanner.NEAR_PIVOT_PCT)
    out["bb_squeeze"] = out["bb_width_pctile"] <= scanner.BB_PCTILE
    out["atr_contract"] = out["atr_ratio"] < scanner.ATR_RATIO
    out["vol_dryup"] = out["vol_10"] < out["vol_50"]
    out["passes_filters"] = ((out["bars"] >= scanner.MIN_BARS) & (c >= scanner.MIN_PRICE)
                             & (c <= scanner.MAX_PRICE) & (out["dollar_vol"] >= scanner.MIN_DOLLAR_VOLUME))
    out["required_ok"] = out["uptrend"] & out["near_pivot"] & out["passes_filters"]
    out["checks"] = out[list(scanner.CONTRACTION)].sum(axis=1).astype(int)
    out["score"] = sum(out[k].astype(int) * w for k, w in scanner.WEIGHTS.items())
    return out


# ------------------------------------------------------------ simulation
def simulate(sym, df, ind, start_date, mode="target"):
    """Walk one stock day by day under one exit rule. Returns a list of trade dicts."""
    cfg = EXIT_MODES[mode]
    trades = []
    o, h, l, c = (df[k].to_numpy() for k in ("open", "high", "low", "close"))
    trail = df["close"].rolling(cfg["trail_sma"]).mean().to_numpy() if cfg.get("trail_sma") else None
    dates = df.index
    req = ind["required_ok"].to_numpy()
    n = len(df)
    in_base = np.zeros(n, dtype=int)            # consecutive days in a base, like Nights on List
    for i in range(n):
        in_base[i] = in_base[i - 1] + 1 if req[i] and i > 0 else int(req[i])

    i = 0
    while i < n - 1:
        if not req[i] or dates[i] < start_date:
            i += 1
            continue
        row = ind.iloc[i]
        entry_px = row["pivot"]

        # buy-stop at the pivot for the next ENTRY_WINDOW days
        fill_i, fill = None, None
        for j in range(i + 1, min(i + 1 + ENTRY_WINDOW, n)):
            if h[j] >= entry_px:
                fill_i, fill = j, max(o[j], entry_px)
                break
        if fill_i is None:
            i += 1
            continue
        r = {"close": row["close"], "pivot": row["pivot"], "base_low": row["base_low"],
             "atr14": row["atr14"], "breakout": False}
        lv = levels.compute(df.iloc[: i + 1], r)
        stop = lv["stop"]
        risk = fill - stop
        if risk <= 0:
            i = fill_i + 1
            continue
        if cfg["target"] == "levels":
            target = lv["target"]
        elif cfg["target"]:
            target = fill + cfg["target"] * risk
        else:
            target = np.inf
        max_hold = cfg["max_hold"]

        exit_i, exit_px, how = None, None, None
        best, active, cur_stop = fill, False, stop
        for k in range(fill_i, min(fill_i + max_hold, n)):
            lo_ok = l[k] <= cur_stop
            hi_ok = h[k] >= target
            if k == fill_i:
                # entry day: the order of the high and low is unknown. Count the stop
                # if the low reached it (conservative); count the target only if the
                # stock also closed at or above it.
                hi_ok = hi_ok and c[k] >= target and not lo_ok
            best = max(best, h[k])
            if lo_ok:
                exit_i, exit_px = k, min(o[k], cur_stop) if k > fill_i else cur_stop
                how = "stop" if cur_stop < fill else "breakeven"
                break
            if hi_ok:
                exit_i, exit_px, how = k, max(o[k], target) if k > fill_i else target, "target"
                break
            if trail is not None:
                if active and c[k] < trail[k]:
                    exit_i, exit_px, how = k, c[k], "trail"
                    break
                if not active and h[k] >= fill + risk:
                    active, cur_stop = True, max(cur_stop, fill)   # breakeven from tomorrow
        if exit_i is None:
            exit_i = min(fill_i + max_hold - 1, n - 1)
            exit_px, how = c[exit_i], "time" if exit_i < n - 1 else "open"

        trades.append({
            "Symbol": sym, "SignalDate": dates[i].date(), "EntryDate": dates[fill_i].date(),
            "ExitDate": dates[exit_i].date(), "Exit": how,
            "Score": int(row["score"]), "Checks": int(row["checks"]),
            "Status": "Setup" if row["checks"] >= scanner.MIN_CHECKS else "Near miss",
            "BBSqueeze": bool(row["bb_squeeze"]), "ATRContract": bool(row["atr_contract"]),
            "VolDryup": bool(row["vol_dryup"]), "PctBelowPivot": round(row["pct_below_pivot"], 2),
            "DaysInBase": int(in_base[i]),
            "ExitRule": mode,
            "Entry": round(fill, 2), "Stop": round(stop, 2),
            "Target": round(target, 2) if np.isfinite(target) else None,
            "PlannedRR": round(lv["reward_risk"], 2) if lv["reward_risk"] == lv["reward_risk"] else None,
            "RROk": bool(lv["rr_ok"]),
            "ExitPrice": round(exit_px, 2),
            "R": round((exit_px - fill) / risk, 3),
            "ReturnPct": round((exit_px - fill) / fill * 100, 2),
            "MaxR": round((best - fill) / risk, 2),
            "HoldDays": int(exit_i - fill_i + 1),
        })
        i = exit_i + 1          # one trade at a time per stock
    return trades


# ------------------------------------------------------------ reporting
def summarize(t, by=None):
    def agg(g):
        r = g["R"]
        pos, neg = r[r > 0].sum(), -r[r < 0].sum()
        return pd.Series({
            "Trades": len(g),
            "Win %": (r > 0).mean() * 100,
            "Avg R": r.mean(),
            "Median R": r.median(),
            "Profit factor": pos / neg if neg > 0 else np.nan,
            "Avg return %": g["ReturnPct"].mean(),
            "Target hit %": (g["Exit"] == "target").mean() * 100,
            "Stopped %": (g["Exit"] == "stop").mean() * 100,
            "Reached 2R %": (g["MaxR"] >= 2).mean() * 100,
            "Avg hold days": g["HoldDays"].mean(),
        })
    if by is None:
        return agg(t).to_frame("All trades").T
    return t.groupby(by, observed=True).apply(agg)


def build_report(t):
    t = t.copy()
    t["Year"] = pd.to_datetime(t["EntryDate"]).dt.year
    t["Score band"] = pd.cut(t["Score"], [0, 79, 89, 100], labels=["Under 80", "80s", "90-100"])
    has_rs = t["RSRating"].notna().sum() >= 50
    if has_rs:
        t["RS quintile"] = pd.qcut(t["RSRating"].rank(method="first"), 5,
                                   labels=["1 (weakest)", "2", "3", "4", "5 (strongest)"])
    t["Days in base"] = pd.cut(t["DaysInBase"], [0, 1, 5, 10, 20, 10_000],
                               labels=["1", "2-5", "6-10", "11-20", "21+"])
    t["Market"] = np.where(t["SPYAbove50"], "SPY above 50-day", "SPY below 50-day")
    t["Planned R:R"] = np.where(t["RROk"], "2:1 or better", "under 2:1")
    setups = t[t["Status"] == "Setup"]
    sheets = {
        "Summary": pd.concat([summarize(setups).rename(index={"All trades": "Setups"}),
                              summarize(t[t["Status"] == "Near miss"]).rename(index={"All trades": "Near misses"})]),
        "By status": summarize(t, "Status"),
        "By checks": summarize(t, "Checks"),
        "By score": summarize(setups, "Score band"),
        "By component": pd.concat({
            f"{name} = {flag}": summarize(setups[setups[name] == flag]).iloc[0]
            for name in ("BBSqueeze", "ATRContract", "VolDryup") for flag in (True, False)
            if (setups[name] == flag).any()}, axis=1).T,
        "By RS": summarize(setups, "RS quintile") if has_rs else pd.DataFrame(),
        "By market": summarize(setups, "Market"),
        "By planned RR": summarize(setups, "Planned R:R"),
        "By days in base": summarize(setups, "Days in base"),
        "By year": summarize(setups, "Year"),
    }
    return sheets


def write_report(sheets, path, meta):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in sheets.items():
            if df is None or df.empty:
                continue
            df = df.round(2)
            df.to_excel(xw, sheet_name=name[:31], startrow=3)
            ws = xw.sheets[name[:31]]
            ws["A1"] = f"Backtest: {name}"
            ws["A1"].font = Font(name="Arial", size=14, bold=True)
            ws["A2"] = meta
            ws["A2"].font = Font(name="Arial", size=9, italic=True, color="555555")
            for row in ws.iter_rows(min_row=4):
                for cell in row:
                    cell.font = Font(name="Arial", size=10, bold=cell.row == 4 or cell.column == 1,
                                     color="FFFFFF" if cell.row == 4 else "000000")
                    if cell.row == 4:
                        cell.fill = PatternFill("solid", start_color="1F4E78")
                        cell.alignment = Alignment(horizontal="center", wrap_text=True)
            ws.column_dimensions["A"].width = 24
            for c in range(2, df.shape[1] + 2):
                ws.column_dimensions[get_column_letter(c)].width = 13
            ws.row_dimensions[4].height = 30
        notes = pd.DataFrame({"How to read": [
            "Avg R is the key number: average result in units of the planned risk. Above 0 = the rules made money on average.",
            "Win % alone can mislead; a 40% win rate with winners twice the size of losers is profitable.",
            "Profit factor = total R won / total R lost. Above 1.5 is solid for a swing system.",
            "Compare rows within a sheet (for example Setups vs Near misses, or RS quintiles), not across market periods.",
            "Survivorship bias: only today's index members are tested, which flatters results. Treat this as a ranking tool for the rules, not a forecast of returns.",
            "Earnings gaps are included but cannot be separated out; no historical earnings dates.",
            "Entry = buy-stop at the pivot for 5 days; exit at stop, target, or after 20 trading days. Same-day stop and target counts as a stop.",
        ]})
        notes.to_excel(xw, sheet_name="How to Read", index=False)
        xw.sheets["How to Read"].column_dimensions["A"].width = 130


def prepare(years=None):
    """Load cached history, compute indicators and RS ratings once."""
    syms = load_symbols()
    spy = load_history("SPY")
    if spy is None:
        sys.exit("No cached history. Run: py backtest.py --download")
    spy_close = spy["close"]
    spy_up = spy_close > spy_close.rolling(50).mean()
    last = spy.index.max()
    start = spy.index.min() + pd.Timedelta(days=400) if years is None \
        else max(spy.index.min() + pd.Timedelta(days=400), last - pd.Timedelta(days=int(years * 365.25)))

    print(f"Computing indicators for {len(syms)} stocks...")
    frames, comp = {}, {}
    for s in syms:
        df = load_history(s)
        if df is None or len(df) < scanner.MIN_BARS:
            continue
        ind = indicator_frame(df)
        frames[s] = (df, ind)
        j = pd.concat([df["close"].rename("s"), spy_close.rename("m")], axis=1, join="inner")
        rel = {}
        for key, n in rs.LOOKBACKS.items():
            rel[key] = (1 + j["s"].pct_change(n)) / (1 + j["m"].pct_change(n)) - 1
        comp[s] = rs.WEIGHTS["rs_3m"] * rel["rs_3m"] + rs.WEIGHTS["rs_6m"] * rel["rs_6m"]
    if not frames:
        sys.exit("No usable history found. Run: py backtest.py --download")
    rating = (pd.DataFrame(comp).rank(axis=1, pct=True) * 98 + 1).round()
    return {"frames": frames, "rating": rating, "spy_up": spy_up, "start": start, "last": last}


def simulate_all(ctx, mode):
    trades = []
    for s, (df, ind) in ctx["frames"].items():
        trades += simulate(s, df, ind, ctx["start"], mode)
    if not trades:
        sys.exit("No trades were generated. Check the cached history.")
    t = pd.DataFrame(trades)
    sd = pd.to_datetime(t["SignalDate"])
    t["SPYAbove50"] = ctx["spy_up"].reindex(sd).fillna(False).astype(bool).to_numpy()
    rating = ctx["rating"]
    t["RSRating"] = [rating.at[d, s] if d in rating.index and s in rating.columns else np.nan
                     for d, s in zip(sd, t["Symbol"])]
    return t


def _save_report(writer, sheets, path, meta):
    try:
        writer(sheets, path, meta)
        return path
    except PermissionError:
        alt = path.with_name(f"{path.stem} ({dt.datetime.now():%H%M}).xlsx")
        writer(sheets, alt, meta)
        return alt


def run(years=None):
    ctx = prepare(years)
    print(f"Simulating trades from {ctx['start'].date()} to {ctx['last'].date()}...")
    t0 = time.time()
    t = simulate_all(ctx, "target")
    TRADES_CSV.parent.mkdir(exist_ok=True)
    t.to_csv(TRADES_CSV, index=False)
    meta = (f"{len(ctx['frames'])} stocks, signals {ctx['start'].date()} to {ctx['last'].date()}, "
            f"{len(t)} simulated trades. Today's index members only (survivorship bias). "
            f"Generated {dt.date.today()}.")
    sheets = build_report(t)
    where = _save_report(write_report, sheets, REPORT, meta)
    s = sheets["Summary"].round(2)
    print(f"\n{meta}\nSimulated in {time.time() - t0:.0f}s\n")
    print(s[["Trades", "Win %", "Avg R", "Profit factor", "Target hit %", "Avg hold days"]].to_string())
    print(f"\nFull breakdown: {where.name}   Every trade: data/{TRADES_CSV.name}")
    return where


# ------------------------------------------------------------ exit comparison
COMPARE_REPORT = HERE / "Backtest Exit Comparison.xlsx"
COMPARE_CSV = HERE / "data" / "backtest_exit_trades.csv"
RS_TOP = 60          # RS rating cutoff for the "top 40%" filter


def _metrics(g, years):
    r = g["R"].sort_index()
    pos, neg = r[r > 0].sum(), -r[r < 0].sum()
    # max drawdown of cumulative R, trades taken in exit-date order
    cum = g.sort_values("ExitDate")["R"].cumsum()
    dd = (cum - cum.cummax()).min() if len(cum) else np.nan
    return pd.Series({
        "Trades": len(g),
        "Trades / yr": len(g) / years if years else np.nan,
        "Win %": (r > 0).mean() * 100,
        "Avg R": r.mean(),
        "Median R": r.median(),
        "Profit factor": pos / neg if neg > 0 else np.nan,
        "Avg win R": r[r > 0].mean(),
        "Avg loss R": r[r < 0].mean(),
        "Trades 3R+ %": (r >= 3).mean() * 100,
        "Total R / yr": r.sum() / years if years else np.nan,
        "Max drawdown R": dd,
        "Avg hold days": g["HoldDays"].mean(),
    })


FILTERS = {
    "All setups": lambda t: t,
    "RS top 40%": lambda t: t[t["RSRating"] >= RS_TOP],
    "RS top 40% + SPY up": lambda t: t[(t["RSRating"] >= RS_TOP) & t["SPYAbove50"]],
}


def compare(years=None):
    ctx = prepare(years)
    span = (ctx["last"] - ctx["start"]).days / 365.25
    print(f"Comparing exit rules on signals {ctx['start'].date()} to {ctx['last'].date()}...")
    allt, rows, yearly = [], {}, {}
    for mode in EXIT_MODES:
        t0 = time.time()
        t = simulate_all(ctx, mode)
        t = t[t["Status"] == "Setup"]
        allt.append(t)
        for fname, f in FILTERS.items():
            rows[(mode, fname)] = _metrics(f(t), span)
        ft = FILTERS["RS top 40%"](t)
        yearly[mode] = ft.groupby(pd.to_datetime(ft["EntryDate"]).dt.year)["R"].mean()
        print(f"  {mode:<8} {len(t):>6} setup trades  ({time.time() - t0:.0f}s)")
    table = pd.DataFrame(rows).T
    table.index.names = ["Exit rule", "Filter"]
    by_year = pd.DataFrame(yearly)
    by_year.index.name = "Year (RS top 40%, avg R)"
    pd.concat(allt).to_csv(COMPARE_CSV, index=False)

    meta = (f"Setups only, signals {ctx['start'].date()} to {ctx['last'].date()}, "
            f"{len(ctx['frames'])} stocks. Today's index members only (survivorship bias).")
    where = _save_report(write_compare, {"Comparison": table, "By year": by_year}, COMPARE_REPORT, meta)
    show = table[["Trades / yr", "Win %", "Avg R", "Profit factor", "Total R / yr",
                  "Max drawdown R", "Avg hold days"]].round(2)
    pd.set_option("display.width", 200)
    print(f"\n{meta}\n")
    print(show.to_string())
    print(f"\nAvg R by year (RS top 40%):\n{by_year.round(2).to_string()}")
    print(f"\nSaved: {where.name}   Every trade: data/{COMPARE_CSV.name}")
    return where


def write_compare(sheets, path, meta):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in sheets.items():
            df.round(2).to_excel(xw, sheet_name=name, startrow=3)
            ws = xw.sheets[name]
            ws["A1"] = f"Backtest exit comparison: {name}"
            ws["A1"].font = Font(name="Arial", size=14, bold=True)
            ws["A2"] = meta
            ws["A2"].font = Font(name="Arial", size=9, italic=True, color="555555")
            for row in ws.iter_rows(min_row=4):
                for cell in row:
                    cell.font = Font(name="Arial", size=10, bold=cell.row == 4,
                                     color="FFFFFF" if cell.row == 4 else "000000")
                    if cell.row == 4:
                        cell.fill = PatternFill("solid", start_color="1F4E78")
                        cell.alignment = Alignment(horizontal="center", wrap_text=True)
            for c in range(1, df.shape[1] + 3):
                ws.column_dimensions[get_column_letter(c)].width = 14
            ws.column_dimensions["B"].width = 22
            ws.row_dimensions[4].height = 30
        notes = pd.DataFrame({"Exit rules": [
            "target: live rules. Conservative target (measured move capped by the nearest swing high), 20-day time exit.",
            "3R: fixed target at 3x the planned risk, 30-day time exit.",
            "trail10: no target. After +1R the stop moves to breakeven, then exit on a close below the 10-day SMA. 40-day limit.",
            "trail20: same with the 20-day SMA, which gives trades more room. 60-day limit.",
            "All rules: buy-stop at the pivot for 5 days, initial stop = tighter of base low and 2x ATR (min 1 ATR), one trade per stock at a time.",
            "Avg R and Total R / yr are the headline numbers. Max drawdown R is the worst peak-to-trough run of cumulative R.",
            "Survivorship bias flatters every rule equally, so compare rules against each other rather than trusting the absolute numbers.",
        ]})
        notes.to_excel(xw, sheet_name="How to Read", index=False)
        xw.sheets["How to Read"].column_dimensions["A"].width = 130


def main():
    ap = argparse.ArgumentParser(description="Backtest the breakout scanner on daily history")
    ap.add_argument("--download", action="store_true", help="download/refresh cached history first")
    ap.add_argument("--download-years", type=float, default=DOWNLOAD_YEARS)
    ap.add_argument("--years", type=float, help="only test the most recent N years")
    ap.add_argument("--source", help="schwab or yahoo (overrides DATA_SOURCE)")
    ap.add_argument("--open", action="store_true", help="open the report in Excel when done")
    ap.add_argument("--compare", action="store_true", help="compare exit rules side by side")
    args = ap.parse_args()
    if args.download:
        download(args.download_years, args.source)
    report = compare(args.years) if args.compare else run(args.years)
    if args.open and report and os.name == "nt":
        os.startfile(report)


if __name__ == "__main__":
    main()
